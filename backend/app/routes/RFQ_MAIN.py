from flask import Blueprint, render_template, session, redirect, url_for, request, flash, send_file
import io
from utils.calculs import (
    calculer_cout_matiere,
    calculer_cout_buy_part,
    calculer_labour_cost,
    calculer_machine_cost,
    appliquer_overhead,
    appliquer_scrap,
    calculer_sa_cost,
    calculer_profit,
    calculer_cout_emballage,
    calculer_ppap_par_1000,
    arrondir
)
from app.models.db import get_db_connection

# Création du Blueprint pour les routes principales
bp = Blueprint('main', __name__)

# Fonction de vérification de connexion
def login_required():
    if 'user_id' not in session:
        flash('Veuillez vous connecter pour accéder à cette page.', 'warning')
        return False
    return True

# ============================================================
# PAGE D'ACCUEIL / TABLEAU DE BORD
# ============================================================
@bp.route('/')
@bp.route('/dashboard')
def dashboard():
    if not login_required():
        return redirect(url_for('auth.login'))
    
    return redirect(url_for('main.dashboard_projets'))

# ============================================================
# TABLEAU DE BORD - Liste des projets
# ============================================================
@bp.route('/dashboard_projets')
def dashboard_projets():
    if not login_required():
        return redirect(url_for('auth.login'))
    
    user_id = session.get('user_id')
    projets = []
    
    conn = get_db_connection()
    if conn:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT id, nom_projet, nom_piece, fournisseur, date_offre, created_at, total_general 
            FROM projets 
            WHERE user_id = %s 
            ORDER BY created_at DESC
        """, (user_id,))
        projets = cursor.fetchall()
        conn.close()
    
    return render_template('dashboard.html',
                         full_name=session.get('full_name', 'Utilisateur'),
                         projets=projets)

# ============================================================
# CHARGER UN PROJET EXISTANT
# ============================================================
@bp.route('/charger_projet/<int:projet_id>')
def charger_projet(projet_id):
    if not login_required():
        return redirect(url_for('auth.login'))
    
    user_id = session.get('user_id')
    
    conn = get_db_connection()
    if conn:
        cursor = conn.cursor(dictionary=True)
        
        cursor.execute("SELECT * FROM projets WHERE id = %s AND user_id = %s", (projet_id, user_id))
        projet = cursor.fetchone()
        
        if projet:
            session['projet_id'] = projet_id
            
            session['page1_data'] = {
                'fournisseur': projet.get('fournisseur', ''),
                'nom_piece': projet.get('nom_piece', ''),
                'part_number': projet.get('part_number', ''),
                'contact': projet.get('contact', ''),
                'date_offre': projet.get('date_offre', ''),
                'devise': projet.get('devise', 'EUR'),
                'taux_change': projet.get('taux_change', '1.0000'),
                'quantite_vie': projet.get('quantite_vie', ''),
                'quantite_an': projet.get('quantite_an', ''),
                'site_production': projet.get('site_production', '')
            }
            
            # Charger les matières
            cursor.execute("SELECT * FROM matieres WHERE projet_id = %s", (projet_id,))
            matieres = cursor.fetchall()
            session['matieres'] = matieres
            
            # Charger les buy-parts
            cursor.execute("SELECT * FROM buy_parts WHERE projet_id = %s", (projet_id,))
            buy_parts = cursor.fetchall()
            session['buy_parts'] = buy_parts
            
            # Charger les processus
            cursor.execute("SELECT * FROM processus WHERE projet_id = %s", (projet_id,))
            processus = cursor.fetchall()
            for proc in processus:
                proc['parts_cycle'] = proc.get('parts_per_cycle')
                proc['manning'] = proc.get('manning_level')
                proc['machine'] = proc.get('machine_type')
            session['processus'] = processus
            
            # Charger les emballages
            cursor.execute("SELECT * FROM emballage WHERE projet_id = %s", (projet_id,))
            emballages = cursor.fetchall()
            session['emballages'] = emballages
            
            # Charger les paramètres
            cursor.execute("SELECT * FROM parametres WHERE projet_id = %s", (projet_id,))
            parametres = cursor.fetchone()
            if parametres:
                session['page5_data'] = {
                    'overhead_mat': parametres.get('overhead_mat', 5.0),
                    'overhead_prod': parametres.get('overhead_prod', 5.0),
                    'scrap_rate': parametres.get('scrap_rate', 3.0)
                }
                session['page6_data'] = {
                    'sa_mat': parametres.get('sa_mat', 6.0),
                    'sa_buy': parametres.get('sa_buy', 0.0),
                    'sa_prod': parametres.get('sa_prod', 6.0),
                    'profit_mat': parametres.get('profit_mat', 6.0),
                    'profit_buy': parametres.get('profit_buy', 0.0),
                    'profit_prod': parametres.get('profit_prod', 6.0)
                }
            
            # Charger le transport
            cursor.execute("SELECT * FROM transport_ppap WHERE projet_id = %s", (projet_id,))
            transport = cursor.fetchone()
            if transport:
                session['page8_data'] = {
                    'incoterms': transport.get('incoterms', 'DAP'),
                    'cout_transport': transport.get('cout_transport', 0.21),
                    'laboratory_test': transport.get('laboratory_test', 0.15)
                }
            
            # Charger les PPAP (base et supplémentaires)
            cursor.execute("SELECT * FROM ppap_supp_projet WHERE projet_id = %s", (projet_id,))
            ppap_supp = cursor.fetchall()
            session['ppap_supp'] = ppap_supp
            
            # Charger ppap_total et quantite_amortissement depuis la table projets
            session['page9_data'] = {
                'ppap_total': projet.get('ppap_total', 1000),
                'quantite_amortissement': projet.get('quantite_amortissement', 0)
            }
            
            flash('Projet chargé avec succès !', 'success')
            return redirect(url_for('main.page1'))
        else:
            flash('Projet non trouvé.', 'danger')
    
    conn.close()
    return redirect(url_for('main.dashboard_projets'))

# ============================================================
# SUPPRIMER UN PROJET
# ============================================================
@bp.route('/supprimer_projet/<int:projet_id>')
def supprimer_projet(projet_id):
    if not login_required():
        return redirect(url_for('auth.login'))
    
    user_id = session.get('user_id')
    
    conn = get_db_connection()
    if conn:
        cursor = conn.cursor()
        cursor.execute("DELETE FROM projets WHERE id = %s AND user_id = %s", (projet_id, user_id))
        conn.commit()
        conn.close()
        
        flash('Projet supprimé avec succès.', 'success')
    
    return redirect(url_for('main.dashboard_projets'))

# ============================================================
# ROUTES DES 10 PAGES
# ============================================================

@bp.route('/page1', methods=['GET', 'POST'])
def page1():
    if not login_required():
        return redirect(url_for('auth.login'))
    
    user_id = session.get('user_id')
    
    if request.method == 'POST':
        fournisseur = request.form.get('fournisseur')
        nom_piece = request.form.get('nom_piece')
        part_number = request.form.get('part_number')
        contact = request.form.get('contact')
        date_offre = request.form.get('date_offre')
        devise = request.form.get('devise')
        taux_change = request.form.get('taux_change')
        quantite_vie = request.form.get('quantite_vie')
        quantite_an = request.form.get('quantite_an')
        site_production = request.form.get('site_production')
        
        session['page1_data'] = {
            'fournisseur': fournisseur,
            'nom_piece': nom_piece,
            'part_number': part_number,
            'contact': contact,
            'date_offre': date_offre,
            'devise': devise,
            'taux_change': taux_change,
            'quantite_vie': quantite_vie,
            'quantite_an': quantite_an,
            'site_production': site_production
        }
        
        conn = get_db_connection()
        if conn:
            cursor = conn.cursor()
            projet_id = session.get('projet_id')
            
            if projet_id:
                cursor.execute("""
                    UPDATE projets 
                    SET fournisseur = %s, nom_piece = %s, part_number = %s, 
                        contact = %s, date_offre = %s, devise = %s, 
                        taux_change = %s, quantite_vie = %s, quantite_an = %s, 
                        site_production = %s, nom_projet = %s
                    WHERE id = %s AND user_id = %s
                """, (fournisseur, nom_piece, part_number, contact, date_offre, 
                      devise, taux_change, quantite_vie, quantite_an, 
                      site_production, f"Projet - {nom_piece}", projet_id, user_id))
            else:
                cursor.execute("""
                    INSERT INTO projets 
                    (user_id, fournisseur, nom_piece, part_number, contact, 
                     date_offre, devise, taux_change, quantite_vie, quantite_an, 
                     site_production, nom_projet)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """, (user_id, fournisseur, nom_piece, part_number, contact, 
                      date_offre, devise, taux_change, quantite_vie, quantite_an, 
                      site_production, f"Projet - {nom_piece}"))
                
                projet_id = cursor.lastrowid
                session['projet_id'] = projet_id
            
            conn.commit()
            conn.close()
        
        flash('Informations générales sauvegardées !', 'success')
        return redirect(url_for('main.page2'))
    
    data = session.get('page1_data', {})
    return render_template('page1.html', 
                         full_name=session.get('full_name', 'Utilisateur'),
                         data=data)

# ============================================================
# PAGE 2 CORRIGÉE : IGNORE LES LIGNES VIDES
# ============================================================
@bp.route('/page2', methods=['GET', 'POST'])
def page2():
    if not login_required():
        return redirect(url_for('auth.login'))
    
    projet_id = session.get('projet_id')
    
    if request.method == 'POST':
        matieres = []
        for i in range(1, 11):
            spec = request.form.get(f'matiere_{i}_spec')
            if spec and spec.strip():
                matieres.append({
                    'no': i,
                    'spec': spec.strip(),
                    'poids_net': request.form.get(f'matiere_{i}_poids_net'),
                    'poids_brut': request.form.get(f'matiere_{i}_poids_brut'),
                    'taux': request.form.get(f'matiere_{i}_taux')
                })
        
        session['matieres'] = matieres
        
        if projet_id:
            conn = get_db_connection()
            if conn:
                cursor = conn.cursor()
                cursor.execute("DELETE FROM matieres WHERE projet_id = %s", (projet_id,))
                
                for mat in matieres:
                    # Vérifier que la ligne n'est pas vide
                    spec = mat.get('spec', '').strip()
                    if not spec:
                        continue
                    
                    cursor.execute("""
                        INSERT INTO matieres 
                        (projet_id, ligne_no, specification, poids_net, poids_brut, taux_matiere)
                        VALUES (%s, %s, %s, %s, %s, %s)
                    """, (projet_id, mat['no'], spec,
                          float(mat['poids_net']) if mat['poids_net'] else 0,
                          float(mat['poids_brut']) if mat['poids_brut'] else 0,
                          float(mat['taux']) if mat['taux'] else 0))
                
                conn.commit()
                conn.close()
        
        flash('Matières sauvegardées !', 'success')
        return redirect(url_for('main.page3'))
    
    matieres = session.get('matieres', [])
    return render_template('page2.html',
                         full_name=session.get('full_name', 'Utilisateur'),
                         matieres=matieres)

# ============================================================
# PAGE 3 CORRIGÉE : IGNORE LES LIGNES VIDES
# ============================================================
@bp.route('/page3', methods=['GET', 'POST'])
def page3():
    if not login_required():
        return redirect(url_for('auth.login'))
    
    projet_id = session.get('projet_id')
    
    if request.method == 'POST':
        buy_parts = []
        for i in range(1, 11):
            spec = request.form.get(f'buy_{i}_spec')
            if spec and spec.strip():
                buy_parts.append({
                    'no': i,
                    'spec': spec.strip(),
                    'quantite': request.form.get(f'buy_{i}_quantite'),
                    'prix_unitaire': request.form.get(f'buy_{i}_prix'),
                    'fournisseur': request.form.get(f'buy_{i}_fournisseur')
                })
        
        session['buy_parts'] = buy_parts
        
        if projet_id:
            conn = get_db_connection()
            if conn:
                cursor = conn.cursor()
                cursor.execute("DELETE FROM buy_parts WHERE projet_id = %s", (projet_id,))
                
                for bp in buy_parts:
                    spec = bp.get('spec', '').strip()
                    if not spec:
                        continue
                    
                    cursor.execute("""
                        INSERT INTO buy_parts 
                        (projet_id, ligne_no, specification, quantite, prix_unitaire, fournisseur)
                        VALUES (%s, %s, %s, %s, %s, %s)
                    """, (projet_id, bp['no'], spec,
                          float(bp['quantite']) if bp['quantite'] else 0,
                          float(bp['prix_unitaire']) if bp['prix_unitaire'] else 0,
                          bp['fournisseur']))
                
                conn.commit()
                conn.close()
        
        flash('Pièces achetées sauvegardées !', 'success')
        return redirect(url_for('main.page4'))
    
    buy_parts = session.get('buy_parts', [])
    return render_template('page3.html',
                         full_name=session.get('full_name', 'Utilisateur'),
                         buy_parts=buy_parts)

# ============================================================
# PAGE 4 CORRIGÉE : IGNORE LES LIGNES VIDES + process_no
# ============================================================
@bp.route('/page4', methods=['GET', 'POST'])
def page4():
    if not login_required():
        return redirect(url_for('auth.login'))
    
    projet_id = session.get('projet_id')
    
    if request.method == 'POST':
        processus = []
        for i in range(1, 11):
            desc = request.form.get(f'process_{i}_desc')
            if desc and desc.strip():
                processus.append({
                    'no': i,
                    'desc': desc.strip(),
                    'machine': request.form.get(f'process_{i}_machine'),
                    'cycle_time': request.form.get(f'process_{i}_cycle'),
                    'parts_cycle': request.form.get(f'process_{i}_parts'),
                    'manning': request.form.get(f'process_{i}_manning'),
                    'labour_rate': request.form.get(f'process_{i}_labour'),
                    'machine_rate': request.form.get(f'process_{i}_machine_rate')
                })
        
        session['processus'] = processus
        
        if projet_id:
            conn = get_db_connection()
            if conn:
                cursor = conn.cursor()
                cursor.execute("DELETE FROM processus WHERE projet_id = %s", (projet_id,))
                
                for proc in processus:
                    desc = proc.get('desc', '').strip()
                    if not desc:
                        continue
                    
                    cursor.execute("""
                        INSERT INTO processus 
                        (projet_id, process_no, description, machine_type, cycle_time, parts_per_cycle, manning_level, labour_rate, machine_rate)
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                    """, (projet_id, str(proc['no']).zfill(2), desc, proc['machine'],
                          float(proc['cycle_time']) if proc['cycle_time'] else 0,
                          int(proc['parts_cycle']) if proc['parts_cycle'] else 1,
                          float(proc['manning']) if proc['manning'] else 0,
                          float(proc['labour_rate']) if proc['labour_rate'] else 0,
                          float(proc['machine_rate']) if proc['machine_rate'] else 0))
                
                conn.commit()
                conn.close()
        
        flash('Processus sauvegardés !', 'success')
        return redirect(url_for('main.page5'))
    
    processus = session.get('processus', [])
    return render_template('page4.html',
                         full_name=session.get('full_name', 'Utilisateur'),
                         processus=processus)

@bp.route('/page5', methods=['GET', 'POST'])
def page5():
    if not login_required():
        return redirect(url_for('auth.login'))
    
    projet_id = session.get('projet_id')
    
    if request.method == 'POST':
        data = {
            'overhead_mat': request.form.get('overhead_mat'),
            'overhead_prod': request.form.get('overhead_prod'),
            'scrap_rate': request.form.get('scrap_rate')
        }
        session['page5_data'] = data
        
        if projet_id:
            conn = get_db_connection()
            if conn:
                cursor = conn.cursor()
                cursor.execute("""
                    INSERT INTO parametres 
                    (projet_id, overhead_mat, overhead_prod, scrap_rate)
                    VALUES (%s, %s, %s, %s)
                    ON DUPLICATE KEY UPDATE
                    overhead_mat = VALUES(overhead_mat),
                    overhead_prod = VALUES(overhead_prod),
                    scrap_rate = VALUES(scrap_rate)
                """, (projet_id, float(data['overhead_mat']) if data['overhead_mat'] else 5.0,
                      float(data['overhead_prod']) if data['overhead_prod'] else 5.0,
                      float(data['scrap_rate']) if data['scrap_rate'] else 3.0))
                conn.commit()
                conn.close()
        
        flash('Taux sauvegardés !', 'success')
        return redirect(url_for('main.page6'))
    
    data = session.get('page5_data', {})
    return render_template('page5.html',
                         full_name=session.get('full_name', 'Utilisateur'),
                         data=data)

@bp.route('/page6', methods=['GET', 'POST'])
def page6():
    if not login_required():
        return redirect(url_for('auth.login'))
    
    projet_id = session.get('projet_id')
    
    if request.method == 'POST':
        data = {
            'sa_mat': request.form.get('sa_mat'),
            'sa_buy': request.form.get('sa_buy'),
            'sa_prod': request.form.get('sa_prod'),
            'profit_mat': request.form.get('profit_mat'),
            'profit_buy': request.form.get('profit_buy'),
            'profit_prod': request.form.get('profit_prod')
        }
        session['page6_data'] = data
        
        if projet_id:
            conn = get_db_connection()
            if conn:
                cursor = conn.cursor()
                cursor.execute("""
                    INSERT INTO parametres 
                    (projet_id, sa_mat, sa_buy, sa_prod, profit_mat, profit_buy, profit_prod)
                    VALUES (%s, %s, %s, %s, %s, %s, %s)
                    ON DUPLICATE KEY UPDATE
                    sa_mat = VALUES(sa_mat),
                    sa_buy = VALUES(sa_buy),
                    sa_prod = VALUES(sa_prod),
                    profit_mat = VALUES(profit_mat),
                    profit_buy = VALUES(profit_buy),
                    profit_prod = VALUES(profit_prod)
                """, (projet_id,
                      float(data['sa_mat']) if data['sa_mat'] else 6.0,
                      float(data['sa_buy']) if data['sa_buy'] else 0.0,
                      float(data['sa_prod']) if data['sa_prod'] else 6.0,
                      float(data['profit_mat']) if data['profit_mat'] else 6.0,
                      float(data['profit_buy']) if data['profit_buy'] else 0.0,
                      float(data['profit_prod']) if data['profit_prod'] else 6.0))
                conn.commit()
                conn.close()
        
        flash('Marges sauvegardées !', 'success')
        return redirect(url_for('main.page7'))
    
    data = session.get('page6_data', {})
    return render_template('page6.html',
                         full_name=session.get('full_name', 'Utilisateur'),
                         data=data)

# ============================================================
# PAGE 7 CORRIGÉE : IGNORE LES LIGNES VIDES
# ============================================================
@bp.route('/page7', methods=['GET', 'POST'])
def page7():
    if not login_required():
        return redirect(url_for('auth.login'))
    
    projet_id = session.get('projet_id')
    
    if request.method == 'POST':
        emballages = []
        for i in range(1, 11):
            desc = request.form.get(f'emb_{i}_desc')
            if desc and desc.strip():
                emballages.append({
                    'no': i,
                    'desc': desc.strip(),
                    'cout_unite': request.form.get(f'emb_{i}_cout'),
                    'pieces_unite': request.form.get(f'emb_{i}_pieces'),
                    'cycles': request.form.get(f'emb_{i}_cycles')
                })
        
        session['emballages'] = emballages
        
        if projet_id:
            conn = get_db_connection()
            if conn:
                cursor = conn.cursor()
                cursor.execute("DELETE FROM emballage WHERE projet_id = %s", (projet_id,))
                
                for emb in emballages:
                    desc = emb.get('desc', '').strip()
                    if not desc:
                        continue
                    
                    cursor.execute("""
                        INSERT INTO emballage 
                        (projet_id, description, cout_unite, pieces_unite, cycles)
                        VALUES (%s, %s, %s, %s, %s)
                    """, (projet_id, desc,
                          float(emb['cout_unite']) if emb['cout_unite'] else 0,
                          int(emb['pieces_unite']) if emb['pieces_unite'] else 0,
                          int(emb['cycles']) if emb['cycles'] else 1))
                
                conn.commit()
                conn.close()
        
        flash('Emballage sauvegardé !', 'success')
        return redirect(url_for('main.page8'))
    
    emballages = session.get('emballages', [])
    return render_template('page7.html',
                         full_name=session.get('full_name', 'Utilisateur'),
                         emballages=emballages)

@bp.route('/page8', methods=['GET', 'POST'])
def page8():
    if not login_required():
        return redirect(url_for('auth.login'))
    
    projet_id = session.get('projet_id')
    
    if request.method == 'POST':
        data = {
            'incoterms': request.form.get('incoterms'),
            'cout_transport': request.form.get('cout_transport'),
            'laboratory_test': request.form.get('laboratory_test')
        }
        session['page8_data'] = data
        
        if projet_id:
            conn = get_db_connection()
            if conn:
                cursor = conn.cursor()
                cursor.execute("""
                    INSERT INTO transport_ppap 
                    (projet_id, incoterms, cout_transport, laboratory_test)
                    VALUES (%s, %s, %s, %s)
                    ON DUPLICATE KEY UPDATE
                    incoterms = VALUES(incoterms),
                    cout_transport = VALUES(cout_transport),
                    laboratory_test = VALUES(laboratory_test)
                """, (projet_id, data['incoterms'],
                      float(data['cout_transport']) if data['cout_transport'] else 0.21,
                      float(data['laboratory_test']) if data['laboratory_test'] else 0.15))
                conn.commit()
                conn.close()
        
        flash('Transport sauvegardé !', 'success')
        return redirect(url_for('main.page9'))
    
    data = session.get('page8_data', {})
    return render_template('page8.html',
                         full_name=session.get('full_name', 'Utilisateur'),
                         data=data)

# ============================================================
# PAGE 9 CORRIGÉE : SAUVEGARDE PPAP DANS LA BDD
# ============================================================
@bp.route('/page9', methods=['GET', 'POST'])
def page9():
    if not login_required():
        return redirect(url_for('auth.login'))
    
    projet_id = session.get('projet_id')
    page1_data = session.get('page1_data', {})
    quantite_vie = float(page1_data.get('quantite_vie', 2931294))
    quantite_amortissement_calc = quantite_vie * 3
    
    if request.method == 'POST':
        ppap_total = request.form.get('ppap_total')
        data = {
            'ppap_total': ppap_total,
            'quantite_amortissement': quantite_amortissement_calc
        }
        session['page9_data'] = data
        
        # Sauvegarde dans la table projets
        if projet_id:
            conn = get_db_connection()
            if conn:
                cursor = conn.cursor()
                cursor.execute("""
                    UPDATE projets 
                    SET ppap_total = %s, quantite_amortissement = %s
                    WHERE id = %s
                """, (float(ppap_total) if ppap_total else 0, quantite_amortissement_calc, projet_id))
                conn.commit()
                conn.close()
        
        # Traitement des PPAP supplémentaires
        ppap_supp = []
        for i in range(1, 11):
            desc = request.form.get(f'ppap_supp_{i}_desc')
            if desc and desc.strip():
                ppap_supp.append({
                    'no': i,
                    'desc': desc.strip(),
                    'montant': request.form.get(f'ppap_supp_{i}_montant')
                })
        
        session['ppap_supp'] = ppap_supp
        
        # Sauvegarde des PPAP supplémentaires en BDD
        if projet_id:
            conn = get_db_connection()
            if conn:
                cursor = conn.cursor()
                cursor.execute("DELETE FROM ppap_supp_projet WHERE projet_id = %s", (projet_id,))
                for supp in ppap_supp:
                    montant = float(supp['montant']) if supp['montant'] else 0
                    cursor.execute("""
                        INSERT INTO ppap_supp_projet (projet_id, ligne_no, description, montant)
                        VALUES (%s, %s, %s, %s)
                    """, (projet_id, supp['no'], supp['desc'], montant))
                conn.commit()
                conn.close()
        
        flash('PPAP sauvegardé !', 'success')
        return redirect(url_for('main.page10'))
    
    # Récupération des données de session (pour l'affichage)
    data = session.get('page9_data', {})
    data['quantite_amortissement'] = quantite_amortissement_calc
    ppap_supp = session.get('ppap_supp', [])
    
    return render_template('page9.html',
                         full_name=session.get('full_name', 'Utilisateur'),
                         data=data,
                         ppap_supp=ppap_supp)

# ============================================================
# PAGE 10 CORRIGÉE
# ============================================================
@bp.route('/page10')
def page10():
    if not login_required():
        return redirect(url_for('auth.login'))
    
    projet_id = session.get('projet_id')
    page1_data = session.get('page1_data', {})
    quantite_vie = float(page1_data.get('quantite_vie', 2931294))
    
    # --- RECHARGEMENT DES DONNÉES DEPUIS LA BDD ---
    if projet_id:
        conn = get_db_connection()
        if conn:
            cursor = conn.cursor(dictionary=True)
            
            # 1. Matières
            cursor.execute("SELECT * FROM matieres WHERE projet_id = %s ORDER BY ligne_no", (projet_id,))
            matieres = cursor.fetchall()
            session['matieres'] = matieres
            
            # 2. Buy-Parts
            cursor.execute("SELECT * FROM buy_parts WHERE projet_id = %s ORDER BY ligne_no", (projet_id,))
            buy_parts = cursor.fetchall()
            session['buy_parts'] = buy_parts
            
            # 3. Processus
            cursor.execute("SELECT * FROM processus WHERE projet_id = %s ORDER BY id", (projet_id,))
            processus = cursor.fetchall()
            for proc in processus:
                proc['parts_cycle'] = proc.get('parts_per_cycle')
                proc['manning'] = proc.get('manning_level')
                proc['machine'] = proc.get('machine_type')
            session['processus'] = processus
            
            # 4. Paramètres
            cursor.execute("SELECT * FROM parametres WHERE projet_id = %s", (projet_id,))
            params = cursor.fetchone()
            if params:
                page5_data = {
                    'overhead_mat': params.get('overhead_mat', 5.0),
                    'overhead_prod': params.get('overhead_prod', 5.0),
                    'scrap_rate': params.get('scrap_rate', 3.0)
                }
                page6_data = {
                    'sa_mat': params.get('sa_mat', 6.0),
                    'sa_buy': params.get('sa_buy', 0.0),
                    'sa_prod': params.get('sa_prod', 6.0),
                    'profit_mat': params.get('profit_mat', 6.0),
                    'profit_buy': params.get('profit_buy', 0.0),
                    'profit_prod': params.get('profit_prod', 6.0)
                }
                session['page5_data'] = page5_data
                session['page6_data'] = page6_data
            else:
                page5_data = session.get('page5_data', {})
                page6_data = session.get('page6_data', {})
            
            # 5. Emballages
            cursor.execute("SELECT * FROM emballage WHERE projet_id = %s ORDER BY id", (projet_id,))
            emballages = cursor.fetchall()
            session['emballages'] = emballages
            
            # 6. Transport
            cursor.execute("SELECT * FROM transport_ppap WHERE projet_id = %s", (projet_id,))
            transport = cursor.fetchone()
            if transport:
                page8_data = {
                    'incoterms': transport.get('incoterms', 'DAP'),
                    'cout_transport': transport.get('cout_transport', 0.21),
                    'laboratory_test': transport.get('laboratory_test', 0.15)
                }
                session['page8_data'] = page8_data
            else:
                page8_data = session.get('page8_data', {})
            
            # 7. PPAP (depuis projets)
            cursor.execute("SELECT ppap_total, quantite_amortissement FROM projets WHERE id = %s", (projet_id,))
            ppap = cursor.fetchone()
            if ppap:
                page9_data = {
                    'ppap_total': ppap.get('ppap_total', 1000),
                    'quantite_amortissement': ppap.get('quantite_amortissement', quantite_vie * 3)
                }
                session['page9_data'] = page9_data
            else:
                page9_data = session.get('page9_data', {})
            
            # 8. PPAP Supplémentaire
            cursor.execute("SELECT * FROM ppap_supp_projet WHERE projet_id = %s ORDER BY ligne_no", (projet_id,))
            ppap_supp_list = cursor.fetchall()
            session['ppap_supp'] = ppap_supp_list
            
            conn.close()
    else:
        # Utiliser les données de session
        matieres = session.get('matieres', [])
        buy_parts = session.get('buy_parts', [])
        processus = session.get('processus', [])
        page5_data = session.get('page5_data', {})
        page6_data = session.get('page6_data', {})
        emballages = session.get('emballages', [])
        page8_data = session.get('page8_data', {})
        page9_data = session.get('page9_data', {})
        ppap_supp_list = session.get('ppap_supp', [])
    
    # ============================================================
    # CALCULS (inchangés)
    # ============================================================
    
    overhead_mat = float(page5_data.get('overhead_mat', 5.0))
    overhead_prod = float(page5_data.get('overhead_prod', 5.0))
    scrap_rate = float(page5_data.get('scrap_rate', 3.0))
    
    sa_mat = float(page6_data.get('sa_mat', 6.0))
    sa_buy = float(page6_data.get('sa_buy', 0.0))
    sa_prod = float(page6_data.get('sa_prod', 6.0))
    profit_mat = float(page6_data.get('profit_mat', 6.0))
    profit_buy = float(page6_data.get('profit_buy', 0.0))
    profit_prod = float(page6_data.get('profit_prod', 6.0))
    
    cout_transport = float(page8_data.get('cout_transport', 0.21))
    laboratory_test = float(page8_data.get('laboratory_test', 0.15))
    
    ppap_total = float(page9_data.get('ppap_total', 1000))
    quantite_amortissement = float(page9_data.get('quantite_amortissement', quantite_vie * 3))
    
    total_matiere = 0
    for mat in matieres:
        cout = calculer_cout_matiere(mat.get('poids_brut'), mat.get('taux_matiere'))
        total_matiere += cout
    
    total_buy_parts = 0
    for bp in buy_parts:
        cout = calculer_cout_buy_part(bp.get('quantite'), bp.get('prix_unitaire'))
        total_buy_parts += cout
    
    total_matiere_avec_buy = total_matiere + total_buy_parts
    overhead_mat_cost = appliquer_overhead(total_matiere_avec_buy, overhead_mat)
    total_matiere_final = total_matiere_avec_buy + overhead_mat_cost
    
    total_labour = 0
    total_machine = 0
    for proc in processus:
        labour = calculer_labour_cost(
            proc.get('cycle_time'),
            proc.get('labour_rate'),
            proc.get('manning_level'),
            proc.get('parts_per_cycle')
        )
        machine = calculer_machine_cost(
            proc.get('cycle_time'),
            proc.get('machine_rate'),
            proc.get('parts_per_cycle')
        )
        total_labour += labour
        total_machine += machine
    
    total_production_direct = total_labour + total_machine
    overhead_prod_cost = appliquer_overhead(total_production_direct, overhead_prod)
    total_production_final = total_production_direct + overhead_prod_cost
    
    scrap_matiere = appliquer_scrap(total_matiere_final, scrap_rate) - total_matiere_final
    scrap_production = appliquer_scrap(total_production_final, scrap_rate) - total_production_final
    total_scrap = scrap_matiere + scrap_production
    
    total_matiere_apres_scrap = appliquer_scrap(total_matiere_final, scrap_rate)
    total_production_apres_scrap = appliquer_scrap(total_production_final, scrap_rate)
    
    sa_matiere = calculer_sa_cost(total_matiere_apres_scrap, sa_mat)
    sa_buy = calculer_sa_cost(total_buy_parts, sa_buy)
    sa_production = calculer_sa_cost(total_production_apres_scrap, sa_prod)
    total_sa = sa_matiere + sa_buy + sa_production
    
    profit_matiere = calculer_profit(total_matiere_apres_scrap + sa_matiere, profit_mat)
    profit_buy = calculer_profit(total_buy_parts + sa_buy, profit_buy)
    profit_production = calculer_profit(total_production_apres_scrap + sa_production, profit_prod)
    total_profit = profit_matiere + profit_buy + profit_production
    
    total_emballage = 0
    for emb in emballages:
        cout = calculer_cout_emballage(
            emb.get('cout_unite'),
            emb.get('pieces_unite'),
            emb.get('cycles')
        )
        total_emballage += cout
    
    ppap_base = calculer_ppap_par_1000(ppap_total, quantite_amortissement)
    total_ppap_supp = 0
    for supp in ppap_supp_list:
        montant = float(supp.get('montant', 0)) if supp.get('montant') else 0
        total_ppap_supp += montant
    ppap_supp_par_1000 = calculer_ppap_par_1000(total_ppap_supp, quantite_amortissement)
    total_ppap = ppap_base + ppap_supp_par_1000
    
    total_general = (
        total_matiere_final +
        total_production_final +
        total_scrap +
        total_sa +
        total_profit +
        total_emballage +
        cout_transport +
        laboratory_test +
        total_ppap
    )
    
    # Sauvegarde du total en BDD
    if projet_id:
        conn = get_db_connection()
        if conn:
            cursor = conn.cursor()
            cursor.execute("UPDATE projets SET total_general = %s WHERE id = %s", (total_general, projet_id))
            conn.commit()
            conn.close()
    
    # Résultats
    resultats = {
        'total_matiere': arrondir(total_matiere_final),
        'total_buy_parts': arrondir(total_buy_parts),
        'total_production': arrondir(total_production_final),
        'overhead_mat': arrondir(overhead_mat_cost),
        'overhead_prod': arrondir(overhead_prod_cost),
        'scrap_matiere': arrondir(scrap_matiere),
        'scrap_production': arrondir(scrap_production),
        'total_scrap': arrondir(total_scrap),
        'sa_matiere': arrondir(sa_matiere),
        'sa_buy': arrondir(sa_buy),
        'sa_production': arrondir(sa_production),
        'total_sa': arrondir(total_sa),
        'profit_matiere': arrondir(profit_matiere),
        'profit_buy': arrondir(profit_buy),
        'profit_production': arrondir(profit_production),
        'total_profit': arrondir(total_profit),
        'emballage': arrondir(total_emballage),
        'transport': arrondir(cout_transport),
        'laboratory_test': arrondir(laboratory_test),
        'ppap_base': arrondir(ppap_base),
        'ppap_supp': arrondir(ppap_supp_par_1000),
        'total_ppap': arrondir(total_ppap),
        'total_general': arrondir(total_general)
    }
    
    if total_general > 0:
        resultats['pct_matiere'] = arrondir((total_matiere_final / total_general) * 100)
        resultats['pct_production'] = arrondir((total_production_final / total_general) * 100)
        resultats['pct_marges'] = arrondir(((total_sa + total_profit) / total_general) * 100)
        resultats['pct_emballage'] = arrondir((total_emballage / total_general) * 100)
        resultats['pct_transport'] = arrondir((cout_transport / total_general) * 100)
        resultats['pct_lab'] = arrondir((laboratory_test / total_general) * 100)
        resultats['pct_ppap'] = arrondir((total_ppap / total_general) * 100)
        resultats['pct_scrap'] = arrondir((total_scrap / total_general) * 100)
    else:
        for key in ['pct_matiere', 'pct_production', 'pct_marges', 'pct_emballage',
                    'pct_transport', 'pct_lab', 'pct_ppap', 'pct_scrap']:
            resultats[key] = 0
    
    session['resultats'] = resultats
    
    return render_template('page10.html',
                         full_name=session.get('full_name', 'Utilisateur'),
                         resultats=resultats)


# ============================================================
# ROUTE DE GÉNÉRATION EXCEL
# ============================================================
@bp.route('/generer_excel')
def generer_excel():
    if not login_required():
        return redirect(url_for('auth.login'))
    
    apply_remise = request.args.get('apply_remise', 'false') == 'true'
    
    page1_data = session.get('page1_data', {})
    matieres = session.get('matieres', [])
    buy_parts = session.get('buy_parts', [])
    processus = session.get('processus', [])
    page5_data = session.get('page5_data', {})
    page6_data = session.get('page6_data', {})
    emballages = session.get('emballages', [])
    page8_data = session.get('page8_data', {})
    page9_data = session.get('page9_data', {})
    ppap_supp_list = session.get('ppap_supp', [])
    
    if apply_remise:
        remise_data = session.get('remise_data', {})
    else:
        remise_data = {}
    
    overhead_mat = float(page5_data.get('overhead_mat', 5.0))
    overhead_prod = float(page5_data.get('overhead_prod', 5.0))
    scrap_rate = float(page5_data.get('scrap_rate', 3.0))
    
    sa_mat = float(page6_data.get('sa_mat', 6.0))
    sa_buy = float(page6_data.get('sa_buy', 0.0))
    sa_prod = float(page6_data.get('sa_prod', 6.0))
    profit_mat = float(page6_data.get('profit_mat', 6.0))
    profit_buy = float(page6_data.get('profit_buy', 0.0))
    profit_prod = float(page6_data.get('profit_prod', 6.0))
    
    quantite_vie = float(page1_data.get('quantite_vie', 2931294))
    quantite_amortissement = quantite_vie * 3
    ppap_total = float(page9_data.get('ppap_total', 1000))
    
    total_matiere = 0
    for mat in matieres:
        cout = calculer_cout_matiere(mat.get('poids_brut'), mat.get('taux_matiere'))
        total_matiere += cout
    
    total_buy_parts = 0
    for bp in buy_parts:
        cout = calculer_cout_buy_part(bp.get('quantite'), bp.get('prix_unitaire'))
        total_buy_parts += cout
    
    total_matiere_avec_buy = total_matiere + total_buy_parts
    overhead_mat_cost = appliquer_overhead(total_matiere_avec_buy, overhead_mat)
    total_matiere_final = total_matiere_avec_buy + overhead_mat_cost
    
    total_labour = 0
    total_machine = 0
    for proc in processus:
        labour = calculer_labour_cost(
            proc.get('cycle_time'),
            proc.get('labour_rate'),
            proc.get('manning_level'),
            proc.get('parts_per_cycle')
        )
        machine = calculer_machine_cost(
            proc.get('cycle_time'),
            proc.get('machine_rate'),
            proc.get('parts_per_cycle')
        )
        total_labour += labour
        total_machine += machine
    
    total_production_direct = total_labour + total_machine
    overhead_prod_cost = appliquer_overhead(total_production_direct, overhead_prod)
    total_production_final = total_production_direct + overhead_prod_cost
    
    scrap_matiere = appliquer_scrap(total_matiere_final, scrap_rate) - total_matiere_final
    scrap_production = appliquer_scrap(total_production_final, scrap_rate) - total_production_final
    total_scrap = scrap_matiere + scrap_production
    
    total_matiere_apres_scrap = appliquer_scrap(total_matiere_final, scrap_rate)
    total_production_apres_scrap = appliquer_scrap(total_production_final, scrap_rate)
    
    sa_matiere = calculer_sa_cost(total_matiere_apres_scrap, sa_mat)
    sa_buy = calculer_sa_cost(total_buy_parts, sa_buy)
    sa_production = calculer_sa_cost(total_production_apres_scrap, sa_prod)
    total_sa = sa_matiere + sa_buy + sa_production
    
    profit_matiere = calculer_profit(total_matiere_apres_scrap + sa_matiere, profit_mat)
    profit_buy = calculer_profit(total_buy_parts + sa_buy, profit_buy)
    profit_production = calculer_profit(total_production_apres_scrap + sa_production, profit_prod)
    total_profit = profit_matiere + profit_buy + profit_production
    
    total_emballage = 0
    for emb in emballages:
        cout = calculer_cout_emballage(
            emb.get('cout_unite'),
            emb.get('pieces_unite'),
            emb.get('cycles')
        )
        total_emballage += cout
    
    ppap_base = calculer_ppap_par_1000(ppap_total, quantite_amortissement)
    total_ppap_supp = 0
    for supp in ppap_supp_list:
        montant = float(supp.get('montant', 0)) if supp.get('montant') else 0
        total_ppap_supp += montant
    ppap_supp_par_1000 = calculer_ppap_par_1000(total_ppap_supp, quantite_amortissement)
    total_ppap = ppap_base + ppap_supp_par_1000
    
    cout_transport = float(page8_data.get('cout_transport', 0.21))
    laboratory_test = float(page8_data.get('laboratory_test', 0.15))
    
    total_general_brut = (
        total_matiere_final +
        total_production_final +
        total_scrap +
        total_sa +
        total_profit +
        total_emballage +
        cout_transport +
        laboratory_test +
        total_ppap
    )
    
    total_general = total_general_brut
    if remise_data:
        total_general = remise_data.get('prix_apres', total_general_brut)
    
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Cost Breakdown"
    
    title_fill = PatternFill(start_color="1a3a6b", end_color="1a3a6b", fill_type="solid")
    header_fill = PatternFill(start_color="2a5f8f", end_color="2a5f8f", fill_type="solid")
    result_fill = PatternFill(start_color="e8f0fe", end_color="e8f0fe", fill_type="solid")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    row = 1
    
    ws.merge_cells(f'A{row}:G{row}')
    ws[f'A{row}'] = "KOSTAL COST BREAKDOWN"
    ws[f'A{row}'].font = Font(bold=True, size=18, color="1a3a6b")
    ws[f'A{row}'].alignment = Alignment(horizontal='center')
    row += 1
    
    ws.merge_cells(f'A{row}:G{row}')
    ws[f'A{row}'] = "RFQ Pro - Plasticum"
    ws[f'A{row}'].font = Font(size=10, italic=True)
    ws[f'A{row}'].alignment = Alignment(horizontal='center')
    row += 2
    
    # 1. Informations Générales
    ws.merge_cells(f'A{row}:G{row}')
    ws[f'A{row}'] = "1. INFORMATIONS GENERALES"
    ws[f'A{row}'].font = Font(bold=True, size=12, color="FFFFFF")
    ws[f'A{row}'].fill = title_fill
    row += 1
    
    info_data = [
        ("Fournisseur", page1_data.get('fournisseur', '')),
        ("Nom de la pièce", page1_data.get('nom_piece', '')),
        ("Part Number", page1_data.get('part_number', '')),
        ("Contact", page1_data.get('contact', '')),
        ("Date de l'offre", page1_data.get('date_offre', '')),
        ("Devise", page1_data.get('devise', 'EUR')),
        ("Taux de change", page1_data.get('taux_change', '1.0000')),
        ("Quantité (Lifetime)", page1_data.get('quantite_vie', '')),
        ("Quantité par an", page1_data.get('quantite_an', '')),
        ("Site de production", page1_data.get('site_production', ''))
    ]
    
    for label, value in info_data:
        ws[f'A{row}'] = label + ":"
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'] = value
        row += 1
    
    row += 1
    
    # 2. Matières
    if matieres:
        ws.merge_cells(f'A{row}:G{row}')
        ws[f'A{row}'] = "2. MATIERES PREMIERES"
        ws[f'A{row}'].font = Font(bold=True, size=12, color="FFFFFF")
        ws[f'A{row}'].fill = title_fill
        row += 1
        
        headers = ['N°', 'Spécification', 'Poids Net (g)', 'Poids Brut (g)', 'Taux (€/kg)', 'Coût (€/1000)']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = Font(bold=True, size=10, color="FFFFFF")
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal='center')
        
        row += 1
        total_cout_matiere = 0
        for i, mat in enumerate(matieres, start=1):
            cout = calculer_cout_matiere(mat.get('poids_brut'), mat.get('taux_matiere'))
            total_cout_matiere += cout
            ws.cell(row=row, column=1, value=i).border = border
            ws.cell(row=row, column=2, value=mat.get('specification', '')).border = border
            ws.cell(row=row, column=3, value=float(mat.get('poids_net', 0)) if mat.get('poids_net') else 0).border = border
            ws.cell(row=row, column=4, value=float(mat.get('poids_brut', 0)) if mat.get('poids_brut') else 0).border = border
            ws.cell(row=row, column=5, value=float(mat.get('taux_matiere', 0)) if mat.get('taux_matiere') else 0).border = border
            ws.cell(row=row, column=6, value=arrondir(cout)).border = border
            row += 1
        
        ws.cell(row=row, column=1, value="TOTAL MATIERE").font = Font(bold=True)
        ws.cell(row=row, column=6, value=arrondir(total_cout_matiere)).font = Font(bold=True)
        ws.cell(row=row, column=6).fill = result_fill
        row += 1
        row += 1
    
    # 3. Buy-Parts
    if buy_parts:
        ws.merge_cells(f'A{row}:G{row}')
        ws[f'A{row}'] = "3. PIECES ACHETEES (BUY-PARTS)"
        ws[f'A{row}'].font = Font(bold=True, size=12, color="FFFFFF")
        ws[f'A{row}'].fill = title_fill
        row += 1
        
        headers = ['N°', 'Spécification', 'Quantité (#/pc)', 'Prix Unitaire (€/pc)', 'Fournisseur', 'Coût (€/1000)']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = Font(bold=True, size=10, color="FFFFFF")
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal='center')
        
        row += 1
        total_cout_buy = 0
        for i, bp in enumerate(buy_parts, start=1):
            cout = calculer_cout_buy_part(bp.get('quantite'), bp.get('prix_unitaire'))
            total_cout_buy += cout
            ws.cell(row=row, column=1, value=i).border = border
            ws.cell(row=row, column=2, value=bp.get('specification', '')).border = border
            ws.cell(row=row, column=3, value=float(bp.get('quantite', 0)) if bp.get('quantite') else 0).border = border
            ws.cell(row=row, column=4, value=float(bp.get('prix_unitaire', 0)) if bp.get('prix_unitaire') else 0).border = border
            ws.cell(row=row, column=5, value=bp.get('fournisseur', '')).border = border
            ws.cell(row=row, column=6, value=arrondir(cout)).border = border
            row += 1
        
        ws.cell(row=row, column=1, value="TOTAL BUY-PARTS").font = Font(bold=True)
        ws.cell(row=row, column=6, value=arrondir(total_cout_buy)).font = Font(bold=True)
        ws.cell(row=row, column=6).fill = result_fill
        row += 1
        row += 1
    
    # 4. Processus
    if processus:
        ws.merge_cells(f'A{row}:G{row}')
        ws[f'A{row}'] = "4. PROCESSUS DE PRODUCTION"
        ws[f'A{row}'].font = Font(bold=True, size=12, color="FFFFFF")
        ws[f'A{row}'].fill = title_fill
        row += 1
        
        headers = ['N°', 'Description', 'Machine', 'Cycle (sec)', 'Pièces/Cycle', 'Taux MO (€/h)', 'Taux Mach.(€/h)']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = Font(bold=True, size=10, color="FFFFFF")
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal='center')
        
        row += 1
        for i, proc in enumerate(processus, start=1):
            ws.cell(row=row, column=1, value=i).border = border
            ws.cell(row=row, column=2, value=proc.get('description', '')).border = border
            ws.cell(row=row, column=3, value=proc.get('machine', '')).border = border
            ws.cell(row=row, column=4, value=float(proc.get('cycle_time', 0)) if proc.get('cycle_time') else 0).border = border
            ws.cell(row=row, column=5, value=int(proc.get('parts_per_cycle', 1)) if proc.get('parts_per_cycle') else 1).border = border
            ws.cell(row=row, column=6, value=float(proc.get('labour_rate', 0)) if proc.get('labour_rate') else 0).border = border
            ws.cell(row=row, column=7, value=float(proc.get('machine_rate', 0)) if proc.get('machine_rate') else 0).border = border
            row += 1
        
        row += 1
    
    # 5. Récapitulatif
    ws.merge_cells(f'A{row}:G{row}')
    ws[f'A{row}'] = "5. RECAPITULATIF DES COUTS"
    ws[f'A{row}'].font = Font(bold=True, size=12, color="FFFFFF")
    ws[f'A{row}'].fill = title_fill
    row += 1
    
    recap_data = [
        ("Matière Première (€/1000)", total_cout_matiere if 'total_cout_matiere' in locals() else 0),
        ("Buy-Parts (€/1000)", total_cout_buy if 'total_cout_buy' in locals() else 0),
        ("Overhead Matière (€/1000)", overhead_mat_cost),
        ("Total Matière (€/1000)", total_matiere_final),
        ("", ""),
        ("Labour Cost (€/1000)", total_labour),
        ("Machine Cost (€/1000)", total_machine),
        ("Overhead Production (€/1000)", overhead_prod_cost),
        ("Total Production (€/1000)", total_production_final),
        ("", ""),
        ("Scrap (€/1000)", total_scrap),
        ("", ""),
        ("S&A Matière (€/1000)", sa_matiere),
        ("S&A Buy-Parts (€/1000)", sa_buy),
        ("S&A Production (€/1000)", sa_production),
        ("Total S&A (€/1000)", total_sa),
        ("", ""),
        ("Profit Matière (€/1000)", profit_matiere),
        ("Profit Buy-Parts (€/1000)", profit_buy),
        ("Profit Production (€/1000)", profit_production),
        ("Total Profit (€/1000)", total_profit),
        ("", ""),
        ("Emballage (€/1000)", total_emballage),
        ("Transport (€/1000)", cout_transport),
        ("Laboratory Test (€/1000)", laboratory_test),
        ("PPAP (€/1000)", total_ppap),
    ]
    
    for label, value in recap_data:
        if label == "":
            row += 1
            continue
        ws[f'A{row}'] = label
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'] = arrondir(value) if isinstance(value, (int, float)) else value
        ws[f'B{row}'].fill = result_fill
        row += 1
    
    row += 1
    
    # 6. LTA / Remise
    if remise_data and apply_remise:
        ws.merge_cells(f'A{row}:G{row}')
        ws[f'A{row}'] = "6. LTA / REMISE COMMERCIALE"
        ws[f'A{row}'].font = Font(bold=True, size=12, color="FFFFFF")
        ws[f'A{row}'].fill = PatternFill(start_color="f59e0b", end_color="f59e0b", fill_type="solid")
        row += 1
        
        ws[f'A{row}'] = "Type de remise:"
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'] = "Sur le prix total" if remise_data.get('type_remise') == 'prix_total' else "Sur le prix hors matière"
        row += 1
        
        ws[f'A{row}'] = "Nom:"
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'] = remise_data.get('nom_remise', 'Remise commerciale')
        row += 1
        
        ws[f'A{row}'] = "Pourcentage (%):"
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'] = remise_data.get('pourcentage', 0)
        row += 1
        
        ws[f'A{row}'] = "Prix avant réduction (€/1000):"
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'] = remise_data.get('prix_avant', 0)
        ws[f'B{row}'].fill = result_fill
        row += 1
        
        ws[f'A{row}'] = "Montant de la réduction (€/1000):"
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'] = remise_data.get('montant_reduction', 0)
        ws[f'B{row}'].fill = PatternFill(start_color="fee2e2", end_color="fee2e2", fill_type="solid")
        ws[f'B{row}'].font = Font(color="dc2626")
        row += 1
        
        ws[f'A{row}'] = "Prix après réduction (€/1000):"
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'] = remise_data.get('prix_apres', 0)
        ws[f'B{row}'].fill = PatternFill(start_color="dcfce7", end_color="dcfce7", fill_type="solid")
        ws[f'B{row}'].font = Font(bold=True, color="16a34a")
        row += 1
        
        row += 1
    
    # 7. Total
    ws.merge_cells(f'A{row}:B{row}')
    ws[f'A{row}'] = "TOTAL COST / 1000 PCS"
    ws[f'A{row}'].font = Font(bold=True, size=14, color="FFFFFF")
    ws[f'A{row}'].fill = PatternFill(start_color="1a3a6b", end_color="1a3a6b", fill_type="solid")
    ws[f'A{row}'].alignment = Alignment(horizontal='center')
    
    ws.merge_cells(f'C{row}:G{row}')
    ws[f'C{row}'] = f"{arrondir(total_general)} €"
    ws[f'C{row}'].font = Font(bold=True, size=14, color="FFFFFF")
    ws[f'C{row}'].fill = PatternFill(start_color="1a3a6b", end_color="1a3a6b", fill_type="solid")
    ws[f'C{row}'].alignment = Alignment(horizontal='center')
    row += 2
    
    for col in range(1, 8):
        ws.column_dimensions[get_column_letter(col)].width = 22
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    nom_fichier = f"devis_{page1_data.get('nom_piece', 'projet')}.xlsx"
    
    return send_file(
        output,
        as_attachment=True,
        download_name=nom_fichier,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


# ============================================================
# PAGE 11 : LTA / REMISE
# ============================================================

@bp.route('/page11', methods=['GET', 'POST'])
def page11():
    if not login_required():
        return redirect(url_for('auth.login'))
    
    page1_data = session.get('page1_data', {})
    resultats = session.get('resultats', {})
    
    if not resultats:
        flash('Veuillez d\'abord calculer les coûts sur la page Résumé.', 'warning')
        return redirect(url_for('main.page10'))
    
    prix_total = resultats.get('total_general', 0)
    total_matiere = resultats.get('total_matiere', 0)
    prix_hors_matiere = prix_total - total_matiere
    
    type_remise = None
    pourcentage = 0
    nom_remise = ""
    prix_avant = 0
    montant_reduction = 0
    prix_apres = 0
    
    if request.method == 'POST':
        type_remise = request.form.get('type_remise')
        pourcentage = float(request.form.get('pourcentage', 0))
        nom_remise = request.form.get('nom_remise', 'Remise commerciale')
        
        if type_remise == 'prix_total':
            prix_avant = prix_total
            montant_reduction = prix_total * (pourcentage / 100)
            prix_apres = prix_total - montant_reduction
        elif type_remise == 'hors_matiere':
            prix_avant = prix_hors_matiere
            montant_reduction = prix_hors_matiere * (pourcentage / 100)
            prix_apres = prix_total - montant_reduction
        else:
            flash('Veuillez sélectionner un type de réduction.', 'danger')
            return redirect(url_for('main.page11'))
        
        session['remise_data'] = {
            'type_remise': type_remise,
            'pourcentage': pourcentage,
            'nom_remise': nom_remise,
            'prix_avant': arrondir(prix_avant),
            'montant_reduction': arrondir(montant_reduction),
            'prix_apres': arrondir(prix_apres),
            'prix_total': arrondir(prix_total),
            'prix_hors_matiere': arrondir(prix_hors_matiere)
        }
        
        projet_id = session.get('projet_id')
        if projet_id:
            conn = get_db_connection()
            if conn:
                cursor = conn.cursor()
                cursor.execute("""
                    INSERT INTO remises 
                    (projet_id, type_remise, pourcentage, nom_remise, 
                     prix_avant, montant_reduction, prix_apres, prix_total, prix_hors_matiere)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                """, (projet_id, type_remise, pourcentage, nom_remise,
                      prix_avant, montant_reduction, prix_apres, prix_total, prix_hors_matiere))
                conn.commit()
                conn.close()
        
        flash('Remise appliquée avec succès !', 'success')
        return redirect(url_for('main.page11'))
    
    remise_data = session.get('remise_data', {})
    
    return render_template('page11.html',
                         full_name=session.get('full_name', 'Utilisateur'),
                         prix_total=arrondir(prix_total),
                         prix_hors_matiere=arrondir(prix_hors_matiere),
                         remise_data=remise_data,
                         resultats=resultats)

@bp.route('/reset_remise')
def reset_remise():
    if 'remise_data' in session:
        session.pop('remise_data')
        flash('Remise supprimée.', 'info')
    return redirect(url_for('main.page11'))