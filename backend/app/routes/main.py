from flask import Blueprint, render_template, session, redirect, url_for, request, flash, send_file
import io
import re
from utils.calculs import (
    calculer_cout_matiere,
    calculer_cout_buy_part,
    calculer_labour_cost,
    calculer_machine_cost,
    appliquer_overhead,
    appliquer_scrap,
    calculer_sa_cost,
    calculer_profit,
    calculer_cout_emballage,
    calculer_ppap_par_1000,
    arrondir
)
from app.models.db import get_db_connection

# Création du Blueprint pour les routes principales
bp = Blueprint('main', __name__)

# ============================================================
# FONCTION UTILITAIRE DE CONVERSION SÉCURISÉE
# ============================================================
def safe_float(value, default=0.0):
    """Convertit une valeur en float, retourne default si échec."""
    if value is None:
        return default
    try:
        return float(value)
    except (ValueError, TypeError):
        return default

# Fonction de vérification de connexion
def login_required():
    if 'user_id' not in session:
        flash('Veuillez vous connecter pour accéder à cette page.', 'warning')
        return False
    return True

# ============================================================
# PAGE D'ACCUEIL / TABLEAU DE BORD
# ============================================================
@bp.route('/')
@bp.route('/dashboard')
def dashboard():
    if not login_required():
        return redirect(url_for('auth.login'))
    return redirect(url_for('main.dashboard_projets'))

# ============================================================
# TABLEAU DE BORD - Liste des projets
# ============================================================
@bp.route('/dashboard_projets')
def dashboard_projets():
    if not login_required():
        return redirect(url_for('auth.login'))
    user_id = session.get('user_id')
    projets = []
    conn = get_db_connection()
    if conn:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT id, nom_projet, nom_piece, fournisseur, date_offre, created_at, total_general 
            FROM projets 
            WHERE user_id = %s 
            ORDER BY created_at DESC
        """, (user_id,))
        projets = cursor.fetchall()
        conn.close()
    return render_template('dashboard.html',
                         full_name=session.get('full_name', 'Utilisateur'),
                         projets=projets)

# ============================================================
# CHARGER UN PROJET EXISTANT
# ============================================================
@bp.route('/charger_projet/<int:projet_id>')
def charger_projet(projet_id):
    if not login_required():
        return redirect(url_for('auth.login'))
    user_id = session.get('user_id')
    conn = get_db_connection()
    if conn:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT * FROM projets WHERE id = %s AND user_id = %s", (projet_id, user_id))
        projet = cursor.fetchone()
        if projet:
            # Vider la session
            session.pop('projet_id', None)
            session.pop('matieres', None)
            session.pop('buy_parts', None)
            session.pop('processus', None)
            session.pop('emballages', None)
            session.pop('ppap_supp', None)
            session.pop('page1_data', None)
            session.pop('page5_data', None)
            session.pop('page6_data', None)
            session.pop('page8_data', None)
            session.pop('page9_data', None)
            session.pop('resultats', None)
            session.pop('remise_data', None)

            session['projet_id'] = projet_id

            date_offre = projet.get('date_offre')
            if date_offre:
                date_offre = str(date_offre).replace('/', '-')
            else:
                date_offre = ''

            session['page1_data'] = {
                'fournisseur': projet.get('fournisseur', ''),
                'nom_piece': projet.get('nom_piece', ''),
                'part_number': projet.get('part_number', ''),
                'contact': projet.get('contact', ''),
                'date_offre': date_offre,
                'devise': projet.get('devise', 'EUR'),
                'taux_change': projet.get('taux_change', '1.0000'),
                'quantite_vie': projet.get('quantite_vie', ''),
                'quantite_an': projet.get('quantite_an', ''),
                'site_production': projet.get('site_production', ''),
                'delivery_lot_size': projet.get('delivery_lot_size', '')
            }

            cursor.execute("SELECT * FROM matieres WHERE projet_id = %s", (projet_id,))
            session['matieres'] = cursor.fetchall()

            cursor.execute("SELECT * FROM buy_parts WHERE projet_id = %s", (projet_id,))
            session['buy_parts'] = cursor.fetchall()

            cursor.execute("SELECT * FROM processus WHERE projet_id = %s", (projet_id,))
            processus = cursor.fetchall()
            for proc in processus:
                proc['parts_cycle'] = proc.get('parts_per_cycle')
                proc['manning'] = proc.get('manning_level')
                proc['machine'] = proc.get('machine_type')
            session['processus'] = processus

            cursor.execute("SELECT * FROM emballage WHERE projet_id = %s", (projet_id,))
            session['emballages'] = cursor.fetchall()

            cursor.execute("SELECT * FROM parametres WHERE projet_id = %s", (projet_id,))
            parametres = cursor.fetchone()
            if parametres:
                # ✅ MODIFICATION : ajout des champs scrap manuel
                session['page5_data'] = {
                    'overhead_mat': parametres.get('overhead_mat', 0),
                    'overhead_prod': parametres.get('overhead_prod', 0),
                    'scrap_rate': parametres.get('scrap_rate', 0),
                    'scrap_manuel_check': parametres.get('scrap_manuel_check', 0) == 1,
                    'scrap_manuel_valeur': parametres.get('scrap_manuel_valeur', 0)
                }
                session['page6_data'] = {
                    'sa_mat': parametres.get('sa_mat', 0),
                    'sa_buy': parametres.get('sa_buy', 0),
                    'sa_prod': parametres.get('sa_prod', 0),
                    'profit_mat': parametres.get('profit_mat', 0),
                    'profit_buy': parametres.get('profit_buy', 0),
                    'profit_prod': parametres.get('profit_prod', 0)
                }
            else:
                session['page5_data'] = {}
                session['page6_data'] = {}

            cursor.execute("SELECT * FROM transport_ppap WHERE projet_id = %s", (projet_id,))
            transport = cursor.fetchone()
            if transport:
                session['page8_data'] = {
                    'incoterms': transport.get('incoterms', ''),
                    'cout_transport': transport.get('cout_transport'),
                    'laboratory_test': transport.get('laboratory_test')
                }
            else:
                session['page8_data'] = {}

            cursor.execute("SELECT * FROM ppap_supp_projet WHERE projet_id = %s", (projet_id,))
            session['ppap_supp'] = cursor.fetchall()

            session['page9_data'] = {
                'ppap_total': projet.get('ppap_total', ''),
                'quantite_amortissement': projet.get('quantite_amortissement', '')
            }

            flash('Projet chargé avec succès !', 'success')
            return redirect(url_for('main.page1'))
        else:
            flash('Projet non trouvé.', 'danger')
    if conn:
        conn.close()
    return redirect(url_for('main.dashboard_projets'))

# ============================================================
# SUPPRIMER UN PROJET
# ============================================================
@bp.route('/supprimer_projet/<int:projet_id>')
def supprimer_projet(projet_id):
    if not login_required():
        return redirect(url_for('auth.login'))
    user_id = session.get('user_id')
    conn = get_db_connection()
    if conn:
        cursor = conn.cursor()
        cursor.execute("DELETE FROM projets WHERE id = %s AND user_id = %s", (projet_id, user_id))
        conn.commit()
        conn.close()
        flash('Projet supprimé avec succès.', 'success')
    return redirect(url_for('main.dashboard_projets'))

# ============================================================
# ROUTES DES 10 PAGES
# ============================================================

@bp.route('/page1', methods=['GET', 'POST'])
def page1():
    if not login_required():
        return redirect(url_for('auth.login'))
    user_id = session.get('user_id')

    if request.method == 'POST':
        fournisseur = request.form.get('fournisseur', '').strip()
        nom_piece = request.form.get('nom_piece', '').strip()
        part_number = request.form.get('part_number', '').strip()
        contact = request.form.get('contact', '').strip()
        date_offre = request.form.get('date_offre', '').strip()
        devise = request.form.get('devise', 'EUR')
        taux_change = request.form.get('taux_change', '').strip()
        quantite_vie = request.form.get('quantite_vie', '').strip()
        quantite_an = request.form.get('quantite_an', '').strip()
        site_production = request.form.get('site_production', '').strip()
        delivery_lot_size = request.form.get('delivery_lot_size', '').strip()

        try:
            taux_change = float(taux_change) if taux_change else 1.0
        except ValueError:
            taux_change = 1.0
        try:
            quantite_vie = int(quantite_vie) if quantite_vie else 0
        except ValueError:
            quantite_vie = 0
        try:
            quantite_an = int(quantite_an) if quantite_an else 0
        except ValueError:
            quantite_an = 0
        try:
            delivery_lot_size = float(delivery_lot_size) if delivery_lot_size else 0.0
        except ValueError:
            delivery_lot_size = 0.0

        session['page1_data'] = {
            'fournisseur': fournisseur,
            'nom_piece': nom_piece,
            'part_number': part_number,
            'contact': contact,
            'date_offre': date_offre,
            'devise': devise,
            'taux_change': taux_change,
            'quantite_vie': quantite_vie,
            'quantite_an': quantite_an,
            'site_production': site_production,
            'delivery_lot_size': delivery_lot_size
        }

        conn = get_db_connection()
        if conn:
            cursor = conn.cursor()
            projet_id = session.get('projet_id')
            if projet_id:
                cursor.execute("""
                    UPDATE projets 
                    SET fournisseur = %s, nom_piece = %s, part_number = %s, 
                        contact = %s, date_offre = %s, devise = %s, 
                        taux_change = %s, quantite_vie = %s, quantite_an = %s, 
                        site_production = %s, nom_projet = %s,
                        delivery_lot_size = %s
                    WHERE id = %s AND user_id = %s
                """, (fournisseur, nom_piece, part_number, contact, date_offre,
                      devise, taux_change, quantite_vie, quantite_an,
                      site_production, f"Projet - {nom_piece}",
                      delivery_lot_size, projet_id, user_id))
            else:
                cursor.execute("""
                    INSERT INTO projets 
                    (user_id, fournisseur, nom_piece, part_number, contact, 
                     date_offre, devise, taux_change, quantite_vie, quantite_an, 
                     site_production, nom_projet, ppap_total, quantite_amortissement,
                     delivery_lot_size)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """, (user_id, fournisseur, nom_piece, part_number, contact,
                      date_offre, devise, taux_change, quantite_vie, quantite_an,
                      site_production, f"Projet - {nom_piece}", None, None,
                      delivery_lot_size))
                projet_id = cursor.lastrowid
                session['projet_id'] = projet_id
                flash(f'Nouveau projet créé (ID: {projet_id}) !', 'success')
            conn.commit()
            conn.close()
        flash('Informations générales sauvegardées !', 'success')
        return redirect(url_for('main.page2'))

    data = session.get('page1_data', {})
    return render_template('page1.html', full_name=session.get('full_name', 'Utilisateur'), data=data)

# ============================================================
# PAGE 2 - Matières
# ============================================================
@bp.route('/page2', methods=['GET', 'POST'])
def page2():
    if not login_required():
        return redirect(url_for('auth.login'))
    projet_id = session.get('projet_id')

    if projet_id and not session.get('matieres'):
        conn = get_db_connection()
        if conn:
            cursor = conn.cursor(dictionary=True)
            cursor.execute("SELECT * FROM matieres WHERE projet_id = %s ORDER BY ligne_no", (projet_id,))
            session['matieres'] = cursor.fetchall()
            conn.close()

    if request.method == 'POST':
        matieres = []
        for key, value in request.form.items():
            if key.startswith('matiere_') and key.endswith('_spec'):
                index = key.split('_')[1]
                spec = value.strip()
                if spec:
                    matieres.append({
                        'no': int(index),
                        'spec': spec,
                        'poids_net': request.form.get(f'matiere_{index}_poids_net'),
                        'poids_brut': request.form.get(f'matiere_{index}_poids_brut'),
                        'taux': request.form.get(f'matiere_{index}_taux')
                    })
        session['matieres'] = matieres

        if projet_id:
            conn = get_db_connection()
            if conn:
                cursor = conn.cursor()
                cursor.execute("DELETE FROM matieres WHERE projet_id = %s", (projet_id,))
                for mat in matieres:
                    cursor.execute("""
                        INSERT INTO matieres 
                        (projet_id, ligne_no, specification, poids_net, poids_brut, taux_matiere)
                        VALUES (%s, %s, %s, %s, %s, %s)
                    """, (projet_id, mat['no'], mat['spec'],
                          float(mat['poids_net']) if mat['poids_net'] else 0,
                          float(mat['poids_brut']) if mat['poids_brut'] else 0,
                          float(mat['taux']) if mat['taux'] else 0))
                conn.commit()
                conn.close()
        flash('Matières sauvegardées !', 'success')
        return redirect(url_for('main.page3'))

    matieres = session.get('matieres', [])
    return render_template('page2.html', full_name=session.get('full_name', 'Utilisateur'), matieres=matieres)

# ============================================================
# PAGE 3 - Buy-Parts
# ============================================================
@bp.route('/page3', methods=['GET', 'POST'])
def page3():
    if not login_required():
        return redirect(url_for('auth.login'))
    projet_id = session.get('projet_id')

    if projet_id and not session.get('buy_parts'):
        conn = get_db_connection()
        if conn:
            cursor = conn.cursor(dictionary=True)
            cursor.execute("SELECT * FROM buy_parts WHERE projet_id = %s ORDER BY ligne_no", (projet_id,))
            session['buy_parts'] = cursor.fetchall()
            conn.close()

    if request.method == 'POST':
        buy_parts = []
        for key, value in request.form.items():
            if key.startswith('buy_') and key.endswith('_spec'):
                index = key.split('_')[1]
                spec = value.strip()
                if spec:
                    buy_parts.append({
                        'no': int(index),
                        'spec': spec,
                        'quantite': request.form.get(f'buy_{index}_quantite'),
                        'prix_unitaire': request.form.get(f'buy_{index}_prix'),
                        'fournisseur': request.form.get(f'buy_{index}_fournisseur')
                    })
        session['buy_parts'] = buy_parts

        if projet_id:
            conn = get_db_connection()
            if conn:
                cursor = conn.cursor()
                cursor.execute("DELETE FROM buy_parts WHERE projet_id = %s", (projet_id,))
                for bp in buy_parts:
                    cursor.execute("""
                        INSERT INTO buy_parts 
                        (projet_id, ligne_no, specification, quantite, prix_unitaire, fournisseur)
                        VALUES (%s, %s, %s, %s, %s, %s)
                    """, (projet_id, bp['no'], bp['spec'],
                          float(bp['quantite']) if bp['quantite'] else 0,
                          float(bp['prix_unitaire']) if bp['prix_unitaire'] else 0,
                          bp['fournisseur']))
                conn.commit()
                conn.close()
        flash('Pièces achetées sauvegardées !', 'success')
        return redirect(url_for('main.page4'))

    buy_parts = session.get('buy_parts', [])
    return render_template('page3.html', full_name=session.get('full_name', 'Utilisateur'), buy_parts=buy_parts)

# ============================================================
# PAGE 4 - Processus
# ============================================================
@bp.route('/page4', methods=['GET', 'POST'])
def page4():
    if not login_required():
        return redirect(url_for('auth.login'))
    projet_id = session.get('projet_id')

    if projet_id and not session.get('processus'):
        conn = get_db_connection()
        if conn:
            cursor = conn.cursor(dictionary=True)
            cursor.execute("SELECT * FROM processus WHERE projet_id = %s ORDER BY id", (projet_id,))
            processus = cursor.fetchall()
            for proc in processus:
                proc['parts_cycle'] = proc.get('parts_per_cycle')
                proc['manning'] = proc.get('manning_level')
                proc['machine'] = proc.get('machine_type')
            session['processus'] = processus
            conn.close()

    if request.method == 'POST':
        processus = []
        for key, value in request.form.items():
            if key.startswith('process_') and key.endswith('_desc'):
                index = key.split('_')[1]
                desc = value.strip()
                if desc:
                    processus.append({
                        'no': int(index),
                        'desc': desc,
                        'machine': request.form.get(f'process_{index}_machine'),
                        'cycle_time': request.form.get(f'process_{index}_cycle'),
                        'parts_cycle': request.form.get(f'process_{index}_parts'),
                        'manning': request.form.get(f'process_{index}_manning'),
                        'labour_rate': request.form.get(f'process_{index}_labour'),
                        'machine_rate': request.form.get(f'process_{index}_machine_rate')
                    })
        session['processus'] = processus

        if projet_id:
            conn = get_db_connection()
            if conn:
                cursor = conn.cursor()
                cursor.execute("DELETE FROM processus WHERE projet_id = %s", (projet_id,))
                for proc in processus:
                    cursor.execute("""
                        INSERT INTO processus 
                        (projet_id, process_no, description, machine_type, cycle_time, parts_per_cycle, manning_level, labour_rate, machine_rate)
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                    """, (projet_id, str(proc['no']).zfill(2), proc['desc'], proc['machine'],
                          float(proc['cycle_time']) if proc['cycle_time'] else 0,
                          int(proc['parts_cycle']) if proc['parts_cycle'] else 1,
                          float(proc['manning']) if proc['manning'] else 0,
                          float(proc['labour_rate']) if proc['labour_rate'] else 0,
                          float(proc['machine_rate']) if proc['machine_rate'] else 0))
                conn.commit()
                conn.close()
        flash('Processus sauvegardés !', 'success')
        return redirect(url_for('main.page5'))

    processus = session.get('processus', [])
    return render_template('page4.html', full_name=session.get('full_name', 'Utilisateur'), processus=processus)

# ============================================================
# PAGE 5 - TAUX (✅ MODIFIÉE : scrap manuel)
# ============================================================
@bp.route('/page5', methods=['GET', 'POST'])
def page5():
    if not login_required():
        return redirect(url_for('auth.login'))
    projet_id = session.get('projet_id')

    if request.method == 'POST':
        data = {
            'overhead_mat': request.form.get('overhead_mat'),
            'overhead_prod': request.form.get('overhead_prod'),
            'scrap_rate': request.form.get('scrap_rate'),
            'scrap_manuel_check': request.form.get('scrap_manuel_check') == 'on',
            'scrap_manuel_valeur': request.form.get('scrap_manuel_valeur')
        }
        session['page5_data'] = data

        if projet_id:
            conn = get_db_connection()
            if conn:
                cursor = conn.cursor()
                cursor.execute("""
                    INSERT INTO parametres 
                    (projet_id, overhead_mat, overhead_prod, scrap_rate,
                     scrap_manuel_check, scrap_manuel_valeur)
                    VALUES (%s, %s, %s, %s, %s, %s)
                    ON DUPLICATE KEY UPDATE
                    overhead_mat = VALUES(overhead_mat),
                    overhead_prod = VALUES(overhead_prod),
                    scrap_rate = VALUES(scrap_rate),
                    scrap_manuel_check = VALUES(scrap_manuel_check),
                    scrap_manuel_valeur = VALUES(scrap_manuel_valeur)
                """, (projet_id,
                      float(data['overhead_mat']) if data['overhead_mat'] else 5.0,
                      float(data['overhead_prod']) if data['overhead_prod'] else 5.0,
                      float(data['scrap_rate']) if data['scrap_rate'] else 3.0,
                      1 if data['scrap_manuel_check'] else 0,
                      float(data['scrap_manuel_valeur']) if data['scrap_manuel_valeur'] else 0))
                conn.commit()
                conn.close()
        flash('Taux sauvegardés !', 'success')
        return redirect(url_for('main.page6'))

    data = session.get('page5_data', {})
    return render_template('page5.html', full_name=session.get('full_name', 'Utilisateur'), data=data)

# ============================================================
# PAGE 6 - MARGES
# ============================================================
@bp.route('/page6', methods=['GET', 'POST'])
def page6():
    if not login_required():
        return redirect(url_for('auth.login'))
    projet_id = session.get('projet_id')

    if request.method == 'POST':
        data = {
            'sa_mat': request.form.get('sa_mat'),
            'sa_buy': request.form.get('sa_buy'),
            'sa_prod': request.form.get('sa_prod'),
            'profit_mat': request.form.get('profit_mat'),
            'profit_buy': request.form.get('profit_buy'),
            'profit_prod': request.form.get('profit_prod')
        }
        session['page6_data'] = data

        if projet_id:
            conn = get_db_connection()
            if conn:
                cursor = conn.cursor()
                cursor.execute("""
                    INSERT INTO parametres 
                    (projet_id, sa_mat, sa_buy, sa_prod, profit_mat, profit_buy, profit_prod)
                    VALUES (%s, %s, %s, %s, %s, %s, %s)
                    ON DUPLICATE KEY UPDATE
                    sa_mat = VALUES(sa_mat),
                    sa_buy = VALUES(sa_buy),
                    sa_prod = VALUES(sa_prod),
                    profit_mat = VALUES(profit_mat),
                    profit_buy = VALUES(profit_buy),
                    profit_prod = VALUES(profit_prod)
                """, (projet_id,
                      float(data['sa_mat']) if data['sa_mat'] else 6.0,
                      float(data['sa_buy']) if data['sa_buy'] else 0.0,
                      float(data['sa_prod']) if data['sa_prod'] else 6.0,
                      float(data['profit_mat']) if data['profit_mat'] else 6.0,
                      float(data['profit_buy']) if data['profit_buy'] else 0.0,
                      float(data['profit_prod']) if data['profit_prod'] else 6.0))
                conn.commit()
                conn.close()
        flash('Marges sauvegardées !', 'success')
        return redirect(url_for('main.page7'))

    data = session.get('page6_data', {})
    return render_template('page6.html', full_name=session.get('full_name', 'Utilisateur'), data=data)

# ============================================================
# PAGE 7 - Emballage
# ============================================================
@bp.route('/page7', methods=['GET', 'POST'])
def page7():
    if not login_required():
        return redirect(url_for('auth.login'))
    projet_id = session.get('projet_id')

    if projet_id and not session.get('emballages'):
        conn = get_db_connection()
        if conn:
            cursor = conn.cursor(dictionary=True)
            cursor.execute("SELECT * FROM emballage WHERE projet_id = %s ORDER BY id", (projet_id,))
            session['emballages'] = cursor.fetchall()
            conn.close()

    if request.method == 'POST':
        emballages = []
        for key, value in request.form.items():
            if key.startswith('emb_') and key.endswith('_desc'):
                index = key.split('_')[1]
                desc = value.strip()
                if desc:
                    emballages.append({
                        'no': int(index),
                        'desc': desc,
                        'cout_unite': request.form.get(f'emb_{index}_cout'),
                        'pieces_unite': request.form.get(f'emb_{index}_pieces'),
                        'cycles': request.form.get(f'emb_{index}_cycles')
                    })
        session['emballages'] = emballages

        if projet_id:
            conn = get_db_connection()
            if conn:
                cursor = conn.cursor()
                cursor.execute("DELETE FROM emballage WHERE projet_id = %s", (projet_id,))
                for emb in emballages:
                    cursor.execute("""
                        INSERT INTO emballage 
                        (projet_id, description, cout_unite, pieces_unite, cycles)
                        VALUES (%s, %s, %s, %s, %s)
                    """, (projet_id, emb['desc'],
                          float(emb['cout_unite']) if emb['cout_unite'] else 0,
                          int(emb['pieces_unite']) if emb['pieces_unite'] else 0,
                          int(emb['cycles']) if emb['cycles'] else 1))
                conn.commit()
                conn.close()
        flash('Emballage sauvegardé !', 'success')
        return redirect(url_for('main.page8'))

    emballages = session.get('emballages', [])
    return render_template('page7.html', full_name=session.get('full_name', 'Utilisateur'), emballages=emballages)

# ============================================================
# PAGE 8 - TRANSPORT
# ============================================================
@bp.route('/page8', methods=['GET', 'POST'])
def page8():
    if not login_required():
        return redirect(url_for('auth.login'))
    projet_id = session.get('projet_id')

    if request.method == 'POST':
        incoterms = request.form.get('incoterms')
        cout_transport = request.form.get('cout_transport')
        laboratory_test = request.form.get('laboratory_test')

        if cout_transport is not None and cout_transport.strip() != '':
            try:
                cout_transport_val = float(cout_transport)
            except ValueError:
                cout_transport_val = 0.0
        else:
            cout_transport_val = 0.0

        if laboratory_test is not None and laboratory_test.strip() != '':
            try:
                laboratory_test_val = float(laboratory_test)
            except ValueError:
                laboratory_test_val = 0.0
        else:
            laboratory_test_val = 0.0

        data = {
            'incoterms': incoterms,
            'cout_transport': cout_transport_val,
            'laboratory_test': laboratory_test_val
        }
        session['page8_data'] = data

        if projet_id:
            conn = get_db_connection()
            if conn:
                cursor = conn.cursor()
                cursor.execute("""
                    INSERT INTO transport_ppap 
                    (projet_id, incoterms, cout_transport, laboratory_test)
                    VALUES (%s, %s, %s, %s)
                    ON DUPLICATE KEY UPDATE
                    incoterms = VALUES(incoterms),
                    cout_transport = VALUES(cout_transport),
                    laboratory_test = VALUES(laboratory_test)
                """, (projet_id, incoterms, cout_transport_val, laboratory_test_val))
                conn.commit()
                conn.close()
        flash('Transport sauvegardé !', 'success')
        return redirect(url_for('main.page9'))

    data = session.get('page8_data', {})
    return render_template('page8.html', full_name=session.get('full_name', 'Utilisateur'), data=data)

# ============================================================
# PAGE 9 - PPAP
# ============================================================
@bp.route('/page9', methods=['GET', 'POST'])
def page9():
    if not login_required():
        return redirect(url_for('auth.login'))

    projet_id = session.get('projet_id')
    page1_data = session.get('page1_data', {})
    quantite_an = safe_float(page1_data.get('quantite_an'))

    if quantite_an == 0 and projet_id:
        conn = get_db_connection()
        if conn:
            cursor = conn.cursor(dictionary=True)
            cursor.execute("SELECT quantite_an FROM projets WHERE id = %s", (projet_id,))
            row = cursor.fetchone()
            if row and row.get('quantite_an') is not None:
                quantite_an = safe_float(row['quantite_an'])
                page1_data['quantite_an'] = quantite_an
                session['page1_data'] = page1_data
            conn.close()

    quantite_amortissement_calc = quantite_an * 3 if quantite_an > 0 else 1
    page1_data['quantite_amortissement'] = quantite_amortissement_calc
    session['page1_data'] = page1_data

    if projet_id and not session.get('ppap_supp'):
        conn = get_db_connection()
        if conn:
            cursor = conn.cursor(dictionary=True)
            cursor.execute("SELECT * FROM ppap_supp_projet WHERE projet_id = %s ORDER BY ligne_no", (projet_id,))
            session['ppap_supp'] = cursor.fetchall()
            conn.close()

    if request.method == 'POST':
        ppap_total = request.form.get('ppap_total')
        if ppap_total is not None and ppap_total.strip() != '':
            try:
                ppap_total_val = float(ppap_total)
            except ValueError:
                ppap_total_val = None
        else:
            ppap_total_val = None

        data = {
            'ppap_total': ppap_total_val,
            'quantite_amortissement': quantite_amortissement_calc
        }
        session['page9_data'] = data

        ppap_supp = []
        for key, value in request.form.items():
            if key.startswith('ppap_supp_') and key.endswith('_desc'):
                index = key.split('_')[2]
                desc = value.strip()
                if desc:
                    ppap_supp.append({
                        'no': int(index),
                        'desc': desc,
                        'montant': request.form.get(f'ppap_supp_{index}_montant')
                    })
        session['ppap_supp'] = ppap_supp

        conn = None
        try:
            if projet_id:
                conn = get_db_connection()
                if conn is None:
                    flash('Erreur de connexion à la base de données.', 'danger')
                    return redirect(url_for('main.page9'))
                cursor = conn.cursor()
                cursor.execute("""
                    UPDATE projets 
                    SET ppap_total = %s, quantite_amortissement = %s
                    WHERE id = %s
                """, (ppap_total_val, quantite_amortissement_calc, projet_id))
                conn.commit()
                cursor.execute("DELETE FROM ppap_supp_projet WHERE projet_id = %s", (projet_id,))
                conn.commit()
                for supp in ppap_supp:
                    montant = float(supp['montant']) if supp['montant'] else 0
                    cursor.execute("""
                        INSERT INTO ppap_supp_projet (projet_id, ligne_no, description, montant)
                        VALUES (%s, %s, %s, %s)
                    """, (projet_id, supp['no'], supp['desc'], montant))
                conn.commit()
                cursor.close()
                conn.close()
            flash('PPAP sauvegardé !', 'success')
            return redirect(url_for('main.page10'))
        except Exception as e:
            if conn:
                conn.rollback()
                conn.close()
            flash(f'Erreur lors de la sauvegarde PPAP : {str(e)}', 'danger')
            return redirect(url_for('main.page9'))

    data = session.get('page9_data', {})
    data['quantite_amortissement'] = quantite_amortissement_calc
    if data.get('ppap_total') is None or data.get('ppap_total') == 0:
        data['ppap_total'] = ''
    session['page9_data'] = data
    ppap_supp = session.get('ppap_supp', [])
    return render_template('page9.html', full_name=session.get('full_name', 'Utilisateur'), data=data, ppap_supp=ppap_supp)

# ============================================================
# PAGE 10 - RÉSUMÉ (✅ MODIFIÉE : scrap manuel)
# ============================================================
@bp.route('/page10')
def page10():
    if not login_required():
        return redirect(url_for('auth.login'))

    projet_id = session.get('projet_id')

    # --- FORCER LE RECHARGEMENT DES INFOS GENERALES DEPUIS LA BDD ---
    if projet_id:
        conn = get_db_connection()
        if conn:
            cursor = conn.cursor(dictionary=True)
            cursor.execute("""
                SELECT fournisseur, nom_piece, part_number, contact, date_offre,
                       devise, taux_change, quantite_vie, quantite_an,
                       site_production, delivery_lot_size
                FROM projets WHERE id = %s
            """, (projet_id,))
            projet = cursor.fetchone()
            if projet:
                page1_data = {
                    'fournisseur': projet.get('fournisseur', ''),
                    'nom_piece': projet.get('nom_piece', ''),
                    'part_number': projet.get('part_number', ''),
                    'contact': projet.get('contact', ''),
                    'date_offre': projet.get('date_offre', ''),
                    'devise': projet.get('devise', 'EUR'),
                    'taux_change': projet.get('taux_change', '1.0000'),
                    'quantite_vie': projet.get('quantite_vie', 0),
                    'quantite_an': projet.get('quantite_an', 0),
                    'site_production': projet.get('site_production', ''),
                    'delivery_lot_size': projet.get('delivery_lot_size', 0)
                }
                session['page1_data'] = page1_data
            conn.close()

    page1_data = session.get('page1_data', {})
    quantite_vie = safe_float(page1_data.get('quantite_vie'))
    quantite_an = safe_float(page1_data.get('quantite_an'))

    # --- RECHARGEMENT DES AUTRES DONNÉES DEPUIS LA BDD ---
    if projet_id:
        conn = get_db_connection()
        if conn:
            cursor = conn.cursor(dictionary=True)

            cursor.execute("SELECT * FROM matieres WHERE projet_id = %s ORDER BY ligne_no", (projet_id,))
            matieres = cursor.fetchall()
            session['matieres'] = matieres

            cursor.execute("SELECT * FROM buy_parts WHERE projet_id = %s ORDER BY ligne_no", (projet_id,))
            buy_parts = cursor.fetchall()
            session['buy_parts'] = buy_parts

            cursor.execute("SELECT * FROM processus WHERE projet_id = %s ORDER BY id", (projet_id,))
            processus = cursor.fetchall()
            for proc in processus:
                proc['parts_cycle'] = proc.get('parts_per_cycle')
                proc['manning'] = proc.get('manning_level')
                proc['machine'] = proc.get('machine_type')
            session['processus'] = processus

            cursor.execute("SELECT * FROM parametres WHERE projet_id = %s", (projet_id,))
            params = cursor.fetchone()
            if params:
                page5_data = {
                    'overhead_mat': params.get('overhead_mat', 5.0),
                    'overhead_prod': params.get('overhead_prod', 5.0),
                    'scrap_rate': params.get('scrap_rate', 3.0),
                    'scrap_manuel_check': params.get('scrap_manuel_check', 0) == 1,
                    'scrap_manuel_valeur': params.get('scrap_manuel_valeur', 0)
                }
                page6_data = {
                    'sa_mat': params.get('sa_mat', 6.0),
                    'sa_buy': params.get('sa_buy', 0.0),
                    'sa_prod': params.get('sa_prod', 6.0),
                    'profit_mat': params.get('profit_mat', 6.0),
                    'profit_buy': params.get('profit_buy', 0.0),
                    'profit_prod': params.get('profit_prod', 6.0)
                }
                session['page5_data'] = page5_data
                session['page6_data'] = page6_data
            else:
                page5_data = session.get('page5_data', {})
                page6_data = session.get('page6_data', {})

            cursor.execute("SELECT * FROM emballage WHERE projet_id = %s ORDER BY id", (projet_id,))
            emballages = cursor.fetchall()
            session['emballages'] = emballages

            cursor.execute("SELECT * FROM transport_ppap WHERE projet_id = %s", (projet_id,))
            transport = cursor.fetchone()
            if transport:
                page8_data = {
                    'incoterms': transport.get('incoterms', 'DAP'),
                    'cout_transport': transport.get('cout_transport'),
                    'laboratory_test': transport.get('laboratory_test')
                }
                session['page8_data'] = page8_data
            else:
                page8_data = session.get('page8_data', {})

            cursor.execute("SELECT ppap_total, quantite_amortissement FROM projets WHERE id = %s", (projet_id,))
            ppap = cursor.fetchone()
            if ppap:
                quantite_amortissement = ppap.get('quantite_amortissement')
                if quantite_amortissement is None:
                    quantite_amortissement = quantite_an * 3 if quantite_an else 1
                page9_data = {
                    'ppap_total': ppap.get('ppap_total', 0),
                    'quantite_amortissement': quantite_amortissement
                }
                session['page9_data'] = page9_data
            else:
                page9_data = session.get('page9_data', {})
                if not page9_data.get('quantite_amortissement'):
                    page9_data['quantite_amortissement'] = quantite_an * 3 if quantite_an else 1
                if page9_data.get('ppap_total') is None:
                    page9_data['ppap_total'] = 0

            cursor.execute("SELECT * FROM ppap_supp_projet WHERE projet_id = %s ORDER BY ligne_no", (projet_id,))
            ppap_supp_list = cursor.fetchall()
            session['ppap_supp'] = ppap_supp_list

            conn.close()
    else:
        matieres = session.get('matieres', [])
        buy_parts = session.get('buy_parts', [])
        processus = session.get('processus', [])
        page5_data = session.get('page5_data', {})
        page6_data = session.get('page6_data', {})
        emballages = session.get('emballages', [])
        page8_data = session.get('page8_data', {})
        page9_data = session.get('page9_data', {})
        ppap_supp_list = session.get('ppap_supp', [])
        if not page9_data.get('quantite_amortissement'):
            page9_data['quantite_amortissement'] = quantite_an * 3 if quantite_an else 1
        if page9_data.get('ppap_total') is None:
            page9_data['ppap_total'] = 0

    # ============================================================
    # CALCULS – alignés sur la logique Excel
    # ============================================================

    overhead_mat = safe_float(page5_data.get('overhead_mat', 5.0))
    overhead_prod = safe_float(page5_data.get('overhead_prod', 5.0))
    scrap_rate = safe_float(page5_data.get('scrap_rate', 3.0))

    sa_mat = safe_float(page6_data.get('sa_mat', 6.0))
    sa_buy = safe_float(page6_data.get('sa_buy', 0.0))
    sa_prod = safe_float(page6_data.get('sa_prod', 6.0))
    profit_mat = safe_float(page6_data.get('profit_mat', 6.0))
    profit_buy = safe_float(page6_data.get('profit_buy', 0.0))
    profit_prod = safe_float(page6_data.get('profit_prod', 6.0))

    cout_transport = safe_float(page8_data.get('cout_transport', 0))
    laboratory_test = safe_float(page8_data.get('laboratory_test', 0))

    ppap_total = safe_float(page9_data.get('ppap_total', 0))
    quantite_amortissement = safe_float(page9_data.get('quantite_amortissement', 1))
    if quantite_amortissement == 0:
        quantite_amortissement = 1

    # --- 1. Matière ---
    total_matiere = 0
    for mat in matieres:
        cout = calculer_cout_matiere(mat.get('poids_brut'), mat.get('taux_matiere'))
        total_matiere += cout

    total_buy_parts = 0
    for bp in buy_parts:
        cout = calculer_cout_buy_part(bp.get('quantite'), bp.get('prix_unitaire'))
        total_buy_parts += cout

    total_matiere_avec_buy = total_matiere + total_buy_parts
    overhead_mat_cost = appliquer_overhead(total_matiere_avec_buy, overhead_mat)
    total_matiere_final = total_matiere_avec_buy + overhead_mat_cost

    # --- 2. Production (séparation normal / setup) ---
    delivery_lot_size = safe_float(page1_data.get('delivery_lot_size', 0))
    if delivery_lot_size == 0:
        delivery_lot_size = quantite_an / 10 if quantite_an > 0 else 1000

    total_labour_normal = 0.0
    total_machine_normal = 0.0
    total_setup_cost = 0.0

    for proc in processus:
        description = proc.get('description', '').lower()
        cycle_time = safe_float(proc.get('cycle_time', 0))
        if cycle_time == 0 and 'setup' in description:
            # Setup
            match = re.search(r'(\d+[.,]?\d*)\s*h', description)
            if match:
                duree_setup_heures = float(match.group(1).replace(',', '.'))
            else:
                duree_setup_heures = 2.0
            labour_rate = safe_float(proc.get('labour_rate', 0))
            machine_rate = safe_float(proc.get('machine_rate', 0))
            cout_setup = (duree_setup_heures * (labour_rate + machine_rate) / delivery_lot_size) * 1000 if delivery_lot_size else 0
            total_setup_cost += cout_setup
        else:
            # Processus normal
            labour = calculer_labour_cost(
                proc.get('cycle_time'),
                proc.get('labour_rate'),
                proc.get('manning_level'),
                proc.get('parts_per_cycle')
            )
            machine = calculer_machine_cost(
                proc.get('cycle_time'),
                proc.get('machine_rate'),
                proc.get('parts_per_cycle')
            )
            total_labour_normal += labour
            total_machine_normal += machine

    # Overhead production uniquement sur la partie normale
    total_production_direct_normal = total_labour_normal + total_machine_normal
    overhead_prod_cost = appliquer_overhead(total_production_direct_normal, overhead_prod)
    total_production_normal = total_production_direct_normal + overhead_prod_cost

    # Production totale (incluant setup, sans overhead sur setup)
    total_production_final = total_production_normal + total_setup_cost

    # --- 3. Scrap (✅ MODIFIÉ : gestion du scrap manuel) ---
    scrap_manuel_check = page5_data.get('scrap_manuel_check', False)
    scrap_manuel_valeur = safe_float(page5_data.get('scrap_manuel_valeur', 0))

    if scrap_manuel_check and scrap_manuel_valeur > 0:
        # Scrap saisi manuellement
        total_scrap = scrap_manuel_valeur
        scrap_matiere = 0
        scrap_production = total_scrap
    else:
        # Calcul automatique classique
        scrap_matiere = total_matiere_final * (scrap_rate / 100)
        scrap_production = total_production_final * (scrap_rate / 100)
        total_scrap = scrap_matiere + scrap_production

    # --- 4. Bases après scrap pour les marges ---
    total_matiere_apres_scrap = total_matiere_final + scrap_matiere
    total_production_apres_scrap = total_production_final + scrap_production

    # --- 5. Marges (S&A et Profit) ---
    sa_matiere = calculer_sa_cost(total_matiere_apres_scrap, sa_mat)
    sa_buy = calculer_sa_cost(total_buy_parts, sa_buy)
    sa_production = calculer_sa_cost(total_production_apres_scrap, sa_prod)
    total_sa = sa_matiere + sa_buy + sa_production

    profit_matiere = calculer_profit(total_matiere_apres_scrap + sa_matiere, profit_mat)
    profit_buy = calculer_profit(total_buy_parts + sa_buy, profit_buy)
    profit_production = calculer_profit(total_production_apres_scrap + sa_production, profit_prod)
    total_profit = profit_matiere + profit_buy + profit_production

    # --- 6. Emballage ---
    total_emballage = 0
    for emb in emballages:
        cout = calculer_cout_emballage(emb.get('cout_unite'), emb.get('pieces_unite'), emb.get('cycles'))
        total_emballage += cout

    # --- 7. PPAP ---
    if ppap_total > 0 and quantite_amortissement > 0:
        ppap_base = calculer_ppap_par_1000(ppap_total, quantite_amortissement)
    else:
        ppap_base = 0

    total_ppap_supp = 0
    for supp in ppap_supp_list:
        montant = safe_float(supp.get('montant', 0))
        total_ppap_supp += montant
    ppap_supp_par_1000 = calculer_ppap_par_1000(total_ppap_supp, quantite_amortissement) if quantite_amortissement else 0
    total_ppap = ppap_base + ppap_supp_par_1000

    # --- 8. Total général ---
    total_general = (
        total_matiere_final +
        total_production_final +
        total_scrap +
        total_sa +
        total_profit +
        total_emballage +
        cout_transport +
        laboratory_test +
        total_ppap
    )

    # Sauvegarde du total en BDD
    if projet_id:
        conn = get_db_connection()
        if conn:
            cursor = conn.cursor()
            cursor.execute("UPDATE projets SET total_general = %s WHERE id = %s", (total_general, projet_id))
            conn.commit()
            conn.close()

    # Résultats
    resultats = {
        'total_matiere': arrondir(total_matiere_final),
        'total_buy_parts': arrondir(total_buy_parts),
        'total_production': arrondir(total_production_final),
        'total_production_normal': arrondir(total_production_normal),
        'total_setup_cost': arrondir(total_setup_cost),
        'overhead_mat': arrondir(overhead_mat_cost),
        'overhead_prod': arrondir(overhead_prod_cost),
        'scrap_matiere': arrondir(scrap_matiere),
        'scrap_production': arrondir(scrap_production),
        'total_scrap': arrondir(total_scrap),
        'sa_matiere': arrondir(sa_matiere),
        'sa_buy': arrondir(sa_buy),
        'sa_production': arrondir(sa_production),
        'total_sa': arrondir(total_sa),
        'profit_matiere': arrondir(profit_matiere),
        'profit_buy': arrondir(profit_buy),
        'profit_production': arrondir(profit_production),
        'total_profit': arrondir(total_profit),
        'emballage': arrondir(total_emballage),
        'transport': arrondir(cout_transport),
        'laboratory_test': arrondir(laboratory_test),
        'ppap_base': arrondir(ppap_base),
        'ppap_supp': arrondir(ppap_supp_par_1000),
        'total_ppap': arrondir(total_ppap),
        'total_general': arrondir(total_general),
        'total_labour': arrondir(total_labour_normal),
        'total_machine': arrondir(total_machine_normal),
        'scrap_manuel_check': scrap_manuel_check,
        'scrap_manuel_valeur': scrap_manuel_valeur
    }

    if total_general > 0:
        resultats['pct_matiere'] = arrondir((total_matiere_final / total_general) * 100)
        resultats['pct_production'] = arrondir((total_production_final / total_general) * 100)
        resultats['pct_marges'] = arrondir(((total_sa + total_profit) / total_general) * 100)
        resultats['pct_emballage'] = arrondir((total_emballage / total_general) * 100)
        resultats['pct_transport'] = arrondir((cout_transport / total_general) * 100)
        resultats['pct_lab'] = arrondir((laboratory_test / total_general) * 100)
        resultats['pct_ppap'] = arrondir((total_ppap / total_general) * 100)
        resultats['pct_scrap'] = arrondir((total_scrap / total_general) * 100)
    else:
        for key in ['pct_matiere', 'pct_production', 'pct_marges', 'pct_emballage',
                    'pct_transport', 'pct_lab', 'pct_ppap', 'pct_scrap']:
            resultats[key] = 0

    session['resultats'] = resultats

    return render_template('page10.html',
                         full_name=session.get('full_name', 'Utilisateur'),
                         resultats=resultats)

# ============================================================
# ROUTE DE GÉNÉRATION EXCEL (✅ MODIFIÉE : scrap manuel)
# ============================================================
@bp.route('/generer_excel')
def generer_excel():
    if not login_required():
        return redirect(url_for('auth.login'))

    apply_remise = request.args.get('apply_remise', 'false') == 'true'

    resultats = session.get('resultats', {})
    if not resultats:
        flash('Veuillez d\'abord calculer les coûts sur la page Résumé.', 'warning')
        return redirect(url_for('main.page10'))

    page1_data = session.get('page1_data', {})
    matieres = session.get('matieres', [])
    buy_parts = session.get('buy_parts', [])
    processus = session.get('processus', [])

    total_matiere_final = resultats.get('total_matiere', 0)
    total_buy_parts = resultats.get('total_buy_parts', 0)
    total_production_final = resultats.get('total_production', 0)
    total_scrap = resultats.get('total_scrap', 0)
    total_sa = resultats.get('total_sa', 0)
    total_profit = resultats.get('total_profit', 0)
    total_emballage = resultats.get('emballage', 0)
    cout_transport = resultats.get('transport', 0)
    laboratory_test = resultats.get('laboratory_test', 0)
    total_ppap = resultats.get('total_ppap', 0)
    total_general = resultats.get('total_general', 0)

    overhead_mat_cost = resultats.get('overhead_mat', 0)
    overhead_prod_cost = resultats.get('overhead_prod', 0)
    total_labour = resultats.get('total_labour', 0)
    total_machine = resultats.get('total_machine', 0)

    if apply_remise:
        remise_data = session.get('remise_data', {})
        if remise_data:
            total_general = remise_data.get('prix_apres', total_general)

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Cost Breakdown"

    title_fill = PatternFill(start_color="1a3a6b", end_color="1a3a6b", fill_type="solid")
    header_fill = PatternFill(start_color="2a5f8f", end_color="2a5f8f", fill_type="solid")
    result_fill = PatternFill(start_color="e8f0fe", end_color="e8f0fe", fill_type="solid")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    row = 1

    ws.merge_cells(f'A{row}:G{row}')
    ws[f'A{row}'] = "KOSTAL COST BREAKDOWN"
    ws[f'A{row}'].font = Font(bold=True, size=18, color="1a3a6b")
    ws[f'A{row}'].alignment = Alignment(horizontal='center')
    row += 1

    ws.merge_cells(f'A{row}:G{row}')
    ws[f'A{row}'] = "RFQ Pro - Plasticum"
    ws[f'A{row}'].font = Font(size=10, italic=True)
    ws[f'A{row}'].alignment = Alignment(horizontal='center')
    row += 2

    # 1. Informations Générales
    ws.merge_cells(f'A{row}:G{row}')
    ws[f'A{row}'] = "1. INFORMATIONS GENERALES"
    ws[f'A{row}'].font = Font(bold=True, size=12, color="FFFFFF")
    ws[f'A{row}'].fill = title_fill
    row += 1

    info_data = [
        ("Fournisseur", page1_data.get('fournisseur', '')),
        ("Nom de la pièce", page1_data.get('nom_piece', '')),
        ("Part Number", page1_data.get('part_number', '')),
        ("Contact", page1_data.get('contact', '')),
        ("Date de l'offre", page1_data.get('date_offre', '')),
        ("Devise", page1_data.get('devise', 'EUR')),
        ("Taux de change", page1_data.get('taux_change', '1.0000')),
        ("Quantité (Lifetime)", page1_data.get('quantite_vie', '')),
        ("Quantité par an", page1_data.get('quantite_an', '')),
        ("Site de production", page1_data.get('site_production', ''))
    ]

    for label, value in info_data:
        ws[f'A{row}'] = label + ":"
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'] = value
        row += 1

    row += 1

    # 2. Matières
    if matieres:
        ws.merge_cells(f'A{row}:G{row}')
        ws[f'A{row}'] = "2. MATIERES PREMIERES"
        ws[f'A{row}'].font = Font(bold=True, size=12, color="FFFFFF")
        ws[f'A{row}'].fill = title_fill
        row += 1

        headers = ['N°', 'Spécification', 'Poids Net (g)', 'Poids Brut (g)', 'Taux (€/kg)', 'Coût (€/1000)']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = Font(bold=True, size=10, color="FFFFFF")
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal='center')

        row += 1
        total_cout_matiere = 0
        for i, mat in enumerate(matieres, start=1):
            cout = calculer_cout_matiere(mat.get('poids_brut'), mat.get('taux_matiere'))
            total_cout_matiere += cout
            ws.cell(row=row, column=1, value=i).border = border
            ws.cell(row=row, column=2, value=mat.get('specification', '')).border = border
            ws.cell(row=row, column=3, value=float(mat.get('poids_net', 0)) if mat.get('poids_net') else 0).border = border
            ws.cell(row=row, column=4, value=float(mat.get('poids_brut', 0)) if mat.get('poids_brut') else 0).border = border
            ws.cell(row=row, column=5, value=float(mat.get('taux_matiere', 0)) if mat.get('taux_matiere') else 0).border = border
            ws.cell(row=row, column=6, value=arrondir(cout)).border = border
            row += 1

        ws.cell(row=row, column=1, value="TOTAL MATIERE").font = Font(bold=True)
        ws.cell(row=row, column=6, value=arrondir(total_cout_matiere)).font = Font(bold=True)
        ws.cell(row=row, column=6).fill = result_fill
        row += 1
        row += 1

    # 3. Buy-Parts
    if buy_parts:
        ws.merge_cells(f'A{row}:G{row}')
        ws[f'A{row}'] = "3. PIECES ACHETEES (BUY-PARTS)"
        ws[f'A{row}'].font = Font(bold=True, size=12, color="FFFFFF")
        ws[f'A{row}'].fill = title_fill
        row += 1

        headers = ['N°', 'Spécification', 'Quantité (#/pc)', 'Prix Unitaire (€/pc)', 'Fournisseur', 'Coût (€/1000)']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = Font(bold=True, size=10, color="FFFFFF")
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal='center')

        row += 1
        total_cout_buy = 0
        for i, bp in enumerate(buy_parts, start=1):
            cout = calculer_cout_buy_part(bp.get('quantite'), bp.get('prix_unitaire'))
            total_cout_buy += cout
            ws.cell(row=row, column=1, value=i).border = border
            ws.cell(row=row, column=2, value=bp.get('specification', '')).border = border
            ws.cell(row=row, column=3, value=float(bp.get('quantite', 0)) if bp.get('quantite') else 0).border = border
            ws.cell(row=row, column=4, value=float(bp.get('prix_unitaire', 0)) if bp.get('prix_unitaire') else 0).border = border
            ws.cell(row=row, column=5, value=bp.get('fournisseur', '')).border = border
            ws.cell(row=row, column=6, value=arrondir(cout)).border = border
            row += 1

        ws.cell(row=row, column=1, value="TOTAL BUY-PARTS").font = Font(bold=True)
        ws.cell(row=row, column=6, value=arrondir(total_cout_buy)).font = Font(bold=True)
        ws.cell(row=row, column=6).fill = result_fill
        row += 1
        row += 1

    # 4. Processus
    if processus:
        ws.merge_cells(f'A{row}:G{row}')
        ws[f'A{row}'] = "4. PROCESSUS DE PRODUCTION"
        ws[f'A{row}'].font = Font(bold=True, size=12, color="FFFFFF")
        ws[f'A{row}'].fill = title_fill
        row += 1

        headers = ['N°', 'Description', 'Machine', 'Cycle (sec)', 'Pièces/Cycle', 'Taux MO (€/h)', 'Taux Mach.(€/h)']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = Font(bold=True, size=10, color="FFFFFF")
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal='center')

        row += 1
        for i, proc in enumerate(processus, start=1):
            ws.cell(row=row, column=1, value=i).border = border
            ws.cell(row=row, column=2, value=proc.get('description', '')).border = border
            ws.cell(row=row, column=3, value=proc.get('machine', '')).border = border
            ws.cell(row=row, column=4, value=float(proc.get('cycle_time', 0)) if proc.get('cycle_time') else 0).border = border
            ws.cell(row=row, column=5, value=int(proc.get('parts_per_cycle', 1)) if proc.get('parts_per_cycle') else 1).border = border
            ws.cell(row=row, column=6, value=float(proc.get('labour_rate', 0)) if proc.get('labour_rate') else 0).border = border
            ws.cell(row=row, column=7, value=float(proc.get('machine_rate', 0)) if proc.get('machine_rate') else 0).border = border
            row += 1

        row += 1

    # 5. Récapitulatif
    ws.merge_cells(f'A{row}:G{row}')
    ws[f'A{row}'] = "5. RECAPITULATIF DES COUTS"
    ws[f'A{row}'].font = Font(bold=True, size=12, color="FFFFFF")
    ws[f'A{row}'].fill = title_fill
    row += 1

    # ✅ MODIFICATION : affichage du scrap manuel dans le récapitulatif
    scrap_label = "Scrap (€/1000)"
    if resultats.get('scrap_manuel_check', False):
        scrap_label = "Scrap - saisi manuellement (€/1000)"
    else:
        scrap_label = "Scrap - calcul auto (€/1000)"

    recap_data = [
        ("Matière Première (€/1000)", total_cout_matiere if 'total_cout_matiere' in locals() else 0),
        ("Buy-Parts (€/1000)", total_cout_buy if 'total_cout_buy' in locals() else 0),
        ("Overhead Matière (€/1000)", overhead_mat_cost),
        ("Total Matière (€/1000)", total_matiere_final),
        ("", ""),
        ("Overhead Production (€/1000)", overhead_prod_cost),
        ("Total Production (€/1000)", total_production_final),
        ("", ""),
        (scrap_label, total_scrap),
        ("", ""),
        ("S&A Matière (€/1000)", resultats.get('sa_matiere', 0)),
        ("S&A Buy-Parts (€/1000)", resultats.get('sa_buy', 0)),
        ("S&A Production (€/1000)", resultats.get('sa_production', 0)),
        ("Total S&A (€/1000)", total_sa),
        ("", ""),
        ("Profit Matière (€/1000)", resultats.get('profit_matiere', 0)),
        ("Profit Buy-Parts (€/1000)", resultats.get('profit_buy', 0)),
        ("Profit Production (€/1000)", resultats.get('profit_production', 0)),
        ("Total Profit (€/1000)", total_profit),
        ("", ""),
        ("Emballage (€/1000)", total_emballage),
        ("Transport (€/1000)", cout_transport),
        ("Laboratory Test (€/1000)", laboratory_test),
        ("PPAP (€/1000)", total_ppap),
    ]

    for label, value in recap_data:
        if label == "":
            row += 1
            continue
        ws[f'A{row}'] = label
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'] = arrondir(value) if isinstance(value, (int, float)) else value
        ws[f'B{row}'].fill = result_fill
        row += 1

    row += 1

    # 6. LTA / Remise
    if apply_remise and session.get('remise_data'):
        remise_data = session.get('remise_data', {})
        ws.merge_cells(f'A{row}:G{row}')
        ws[f'A{row}'] = "6. LTA / REMISE COMMERCIALE"
        ws[f'A{row}'].font = Font(bold=True, size=12, color="FFFFFF")
        ws[f'A{row}'].fill = PatternFill(start_color="f59e0b", end_color="f59e0b", fill_type="solid")
        row += 1

        ws[f'A{row}'] = "Type de remise:"
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'] = "Sur le prix total" if remise_data.get('type_remise') == 'prix_total' else "Sur le prix hors matière"
        row += 1

        ws[f'A{row}'] = "Nom:"
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'] = remise_data.get('nom_remise', 'Remise commerciale')
        row += 1

        ws[f'A{row}'] = "Pourcentage (%):"
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'] = remise_data.get('pourcentage', 0)
        row += 1

        ws[f'A{row}'] = "Prix avant réduction (€/1000):"
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'] = remise_data.get('prix_avant', 0)
        ws[f'B{row}'].fill = result_fill
        row += 1

        ws[f'A{row}'] = "Montant de la réduction (€/1000):"
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'] = remise_data.get('montant_reduction', 0)
        ws[f'B{row}'].fill = PatternFill(start_color="fee2e2", end_color="fee2e2", fill_type="solid")
        ws[f'B{row}'].font = Font(color="dc2626")
        row += 1

        ws[f'A{row}'] = "Prix après réduction (€/1000):"
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'] = remise_data.get('prix_apres', 0)
        ws[f'B{row}'].fill = PatternFill(start_color="dcfce7", end_color="dcfce7", fill_type="solid")
        ws[f'B{row}'].font = Font(bold=True, color="16a34a")
        row += 1

        row += 1

    # 7. Total
    ws.merge_cells(f'A{row}:B{row}')
    ws[f'A{row}'] = "TOTAL COST / 1000 PCS"
    ws[f'A{row}'].font = Font(bold=True, size=14, color="FFFFFF")
    ws[f'A{row}'].fill = PatternFill(start_color="1a3a6b", end_color="1a3a6b", fill_type="solid")
    ws[f'A{row}'].alignment = Alignment(horizontal='center')

    ws.merge_cells(f'C{row}:G{row}')
    ws[f'C{row}'] = f"{arrondir(total_general)} €"
    ws[f'C{row}'].font = Font(bold=True, size=14, color="FFFFFF")
    ws[f'C{row}'].fill = PatternFill(start_color="1a3a6b", end_color="1a3a6b", fill_type="solid")
    ws[f'C{row}'].alignment = Alignment(horizontal='center')
    row += 2

    for col in range(1, 8):
        ws.column_dimensions[get_column_letter(col)].width = 22

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    nom_fichier = f"devis_{page1_data.get('nom_piece', 'projet')}.xlsx"

    return send_file(
        output,
        as_attachment=True,
        download_name=nom_fichier,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

# ============================================================
# PAGE 11 : LTA / REMISE
# ============================================================

@bp.route('/page11', methods=['GET', 'POST'])
def page11():
    if not login_required():
        return redirect(url_for('auth.login'))

    page1_data = session.get('page1_data', {})
    resultats = session.get('resultats', {})

    if not resultats:
        flash('Veuillez d\'abord calculer les coûts sur la page Résumé.', 'warning')
        return redirect(url_for('main.page10'))

    prix_total = resultats.get('total_general', 0)
    total_matiere = resultats.get('total_matiere', 0)
    prix_hors_matiere = prix_total - total_matiere

    type_remise = None
    pourcentage = 0
    nom_remise = ""
    prix_avant = 0
    montant_reduction = 0
    prix_apres = 0

    if request.method == 'POST':
        type_remise = request.form.get('type_remise')
        pourcentage = float(request.form.get('pourcentage', 0))
        nom_remise = request.form.get('nom_remise', 'Remise commerciale')

        if type_remise == 'prix_total':
            prix_avant = prix_total
            montant_reduction = prix_total * (pourcentage / 100)
            prix_apres = prix_total - montant_reduction
        elif type_remise == 'hors_matiere':
            prix_avant = prix_hors_matiere
            montant_reduction = prix_hors_matiere * (pourcentage / 100)
            prix_apres = prix_total - montant_reduction
        else:
            flash('Veuillez sélectionner un type de réduction.', 'danger')
            return redirect(url_for('main.page11'))

        session['remise_data'] = {
            'type_remise': type_remise,
            'pourcentage': pourcentage,
            'nom_remise': nom_remise,
            'prix_avant': arrondir(prix_avant),
            'montant_reduction': arrondir(montant_reduction),
            'prix_apres': arrondir(prix_apres),
            'prix_total': arrondir(prix_total),
            'prix_hors_matiere': arrondir(prix_hors_matiere)
        }

        projet_id = session.get('projet_id')
        if projet_id:
            conn = get_db_connection()
            if conn:
                cursor = conn.cursor()
                cursor.execute("""
                    INSERT INTO remises 
                    (projet_id, type_remise, pourcentage, nom_remise, 
                     prix_avant, montant_reduction, prix_apres, prix_total, prix_hors_matiere)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                """, (projet_id, type_remise, pourcentage, nom_remise,
                      prix_avant, montant_reduction, prix_apres, prix_total, prix_hors_matiere))
                conn.commit()
                conn.close()

        flash('Remise appliquée avec succès !', 'success')
        return redirect(url_for('main.page11'))

    remise_data = session.get('remise_data', {})

    return render_template('page11.html',
                         full_name=session.get('full_name', 'Utilisateur'),
                         prix_total=arrondir(prix_total),
                         prix_hors_matiere=arrondir(prix_hors_matiere),
                         remise_data=remise_data,
                         resultats=resultats)

@bp.route('/reset_remise')
def reset_remise():
    if 'remise_data' in session:
        session.pop('remise_data')
        flash('Remise supprimée.', 'info')
    return redirect(url_for('main.page11'))

# ============================================================
# NOUVEAU PROJET - Vide la session avant de créer
# ============================================================
@bp.route('/nouveau_projet')
def nouveau_projet():
    if not login_required():
        return redirect(url_for('auth.login'))

    session.pop('projet_id', None)
    session.pop('matieres', None)
    session.pop('buy_parts', None)
    session.pop('processus', None)
    session.pop('emballages', None)
    session.pop('ppap_supp', None)
    session.pop('page1_data', None)
    session.pop('page5_data', None)
    session.pop('page6_data', None)
    session.pop('page8_data', None)
    session.pop('page9_data', None)
    session.pop('resultats', None)
    session.pop('remise_data', None)

    flash('Nouveau projet créé !', 'success')
    return redirect(url_for('main.page1'))